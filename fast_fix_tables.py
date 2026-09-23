import docx
from lxml import etree
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_table_caption(p, seq_num, desc_text):
    for child in list(p._p):
        tag = child.tag.split('}')[-1]
        if tag not in ('pPr',):
            p._p.remove(child)
            
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:keepNext')) is None:
        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    
    # Bold prefix: "Tabel 4."
    r1 = p.add_run("Tabel 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    # SEQ field (BOLD)
    fld_xml = '''
        <w:fldSimple %s w:instr=" SEQ Tabel \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="24"/>
                </w:rPr>
                <w:t>%d</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), seq_num)
    p._p.append(parse_xml(fld_xml))
    
    # Regular desc (NOT BOLD)
    r2 = p.add_run(f" {desc_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = False

def run_fix():
    doc = docx.Document('SKRIPSI.docx')
    
    target_tables = [
        ("Kebutuhan Perangkat Keras", "Kebutuhan Perangkat Keras"),
        ("Kebutuhan Perangkat Lunak", "Kebutuhan Perangkat Lunak"),
        ("Identifikasi Aktor", "Identifikasi Aktor"),
        ("Pemetaan Kebutuhan Fungsional terhadap Use Case", "Pemetaan Kebutuhan Fungsional terhadap Use Case"),
        ("Deskripsi Use Case", "Deskripsi Use Case"),
        ("Skenario Use Case Login", "Skenario Use Case Login"),
        ("Skenario Use Case Mengelola Data Pengguna", "Skenario Use Case Mengelola Data Pengguna"),
        ("Skenario Use Case Mengelola Data Guru", "Skenario Use Case Mengelola Data Guru"),
        ("Skenario Use Case Mengelola Kriteria dan Bobot", "Skenario Use Case Mengelola Kriteria dan Bobot"),
        ("Skenario Use Case  Mengelola Skala Penilaian", "Skenario Use Case Mengelola Skala Penilaian"),
        ("Skenario Use Case Mengelola Periode Penilaian", "Skenario Use Case Mengelola Periode Penilaian"),
        ("Skenario Use Case  Melakukan Penilaian Kinerja Guru", "Skenario Use Case Melakukan Penilaian Kinerja Guru"),
        ("Skenario Use Case Mencatat Hasil Penilaian", "Skenario Use Case Mencatat Hasil Penilaian"),
        ("Skenario Use Case Mengolah Penilaian dengan MOORA", "Skenario Use Case Mengolah Penilaian dengan MOORA"),
        ("Skenario Use Case  Melihat Hasil Penilaian", "Skenario Use Case Melihat Hasil Penilaian"),
        ("Skenario Use Case  Melihat Pemeringkatan", "Skenario Use Case Melihat Pemeringkatan"),
        ("Skenario Use Case  Monitoring Kinerja Guru", "Skenario Use Case Monitoring Kinerja Guru"),
        ("Skenario Use Case  Melihat Detail Penilaian", "Skenario Use Case Melihat Detail Penilaian"),
        ("Skenario Use Case  Melihat Riwayat Penilaian", "Skenario Use Case Melihat Riwayat Penilaian"),
        ("Skenario Use Case  Mencetak Laporan Penilaian", "Skenario Use Case Mencetak Laporan Penilaian"),
        ("Struktur Tabel Role", "Struktur Tabel Role (tbl_role)"),
        ("Struktur Tabel User", "Struktur Tabel User (tbl_user)"),
        ("Struktur Tabel Guru", "Struktur Tabel Guru (tbl_guru)"),
        ("Struktur Tabel Kriteria", "Struktur Tabel Kriteria (tbl_kriteria)"),
        ("Struktur Tabel Skala Penilaian", "Struktur Tabel Skala Penilaian (tbl_skala_penilaian)"),
        ("Struktur Tabel Periode Penilaian", "Struktur Tabel Periode Penilaian (tbl_periode_penilaian)"),
        ("Struktur Tabel Penilaian", "Struktur Tabel Penilaian (tbl_penilaian)"),
        ("Struktur Tabel Detail Penilaian", "Struktur Tabel Detail Penilaian (tbl_detail_penilaian)"),
        ("Struktur Tabel Hasil MOORA", "Struktur Tabel Hasil MOORA (tbl_hasil_moora)"),
        ("Struktur Tabel Detail MOORA", "Struktur Tabel Detail MOORA (tbl_detail_moora)"),
        ("Struktur Tabel Log Aktivitas", "Struktur Tabel Log Aktivitas (tbl_log_aktivitas)"),
        ("Daftar Kriteria, Bobot Preferensi, dan Sifat Atribut MOORA", "Daftar Kriteria, Bobot Preferensi, dan Sifat Atribut MOORA"),
        ("Daftar Alternatif Guru SMA Al-Ihsan Boarding School", "Daftar Alternatif Guru SMA Al-Ihsan Boarding School"),
        ("Matriks Keputusan (X) Evaluasi Kinerja Guru", "Matriks Keputusan (X) Evaluasi Kinerja Guru"),
        ("Matriks Ternormalisasi (X*) Metode MOORA", "Matriks Ternormalisasi (X*) Metode MOORA"),
        ("Matriks Normalisasi Terbobot (W x X*)", "Matriks Normalisasi Terbobot (W x X*)"),
        ("Hasil Akhir Nilai Preferensi (Yi) dan Peringkat Kinerja Guru (MOORA)", "Hasil Akhir Nilai Preferensi (Yi) dan Peringkat Kinerja Guru (MOORA)"),
    ]
    
    # Process BAB IV paragraphs (index > 700)
    for seq_num, (marker, desc_clean) in enumerate(target_tables, start=1):
        found = False
        for i, p in enumerate(doc.paragraphs):
            if i > 700:
                t = p.text.strip()
                if marker in t:
                    set_table_caption(p, seq_num, desc_clean)
                    print(f"[OK #{seq_num:2d}] Tabel 4.{seq_num:2d} {desc_clean}")
                    found = True
                    break
        if not found:
            print(f"[FAILED #{seq_num:2d}] marker='{marker}'")
            
    doc.save('SKRIPSI.docx')
    print("[SUCCESS] SKRIPSI.docx saved!")

if __name__ == "__main__":
    run_fix()
