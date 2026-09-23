import os
import glob
import docx
from lxml import etree

print("Searching for docx files in directory:")
for f in glob.glob("*.docx"):
    print(f" - {f}")

# Check if there's any file with AD-11 image
for fn in glob.glob("*.docx"):
    if 'SKRIPSI' in fn:
        try:
            d = docx.Document(fn)
            for i, p in enumerate(d.paragraphs):
                if 'AD-11' in p.text:
                    # check next 3 paragraphs for drawing
                    has_img = False
                    for off in range(1, 4):
                        if i + off < len(d.paragraphs):
                            if 'w:drawing' in d.paragraphs[i+off]._p.xml:
                                has_img = True
                    print(f"File {fn}: AD-11 found at P{i}, has_image={has_img}")
        except Exception as e:
            print(f"Error reading {fn}: {e}")
