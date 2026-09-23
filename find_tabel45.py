import docx
from docx.oxml.ns import qn

doc = docx.Document('SKRIPSI.docx')

print("Searching for Tabel 4.5...")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if 'Tabel 4.' in t or 'Tabel 4 ' in t:
        if '5' in t or 'Deskripsi' in t or 'Skala' in t:
            print(f"P{i}: '{t}'")
            for offset in range(-2, 4):
                j = i + offset
                if 0 <= j < len(doc.paragraphs):
                    pj = doc.paragraphs[j]
                    print(f"   P{j} (sb={pj.paragraph_format.space_before}, sa={pj.paragraph_format.space_after}, style={pj.style.name}): '{pj.text[:60]}'")
            print("-" * 50)
