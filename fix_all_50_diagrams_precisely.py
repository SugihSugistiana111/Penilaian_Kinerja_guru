import docx
from lxml import etree
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml, OxmlElement
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_caption(p, seq_num, desc_text):
    # Clear runs
    for child in list(p._p):
        tag = child.tag.split('}')[-1]
        if tag not in ('pPr',):
            p._p.remove(child)
            
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:keepNext')) is None:
        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    
    # Bold prefix: "Gambar 4."
    r1 = p.add_run("Gambar 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    # SEQ field (BOLD)
    fld_xml = '''
        <w:fldSimple %s w:instr=" SEQ Gambar \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="24"/>
                </w:rPr>
                <w:t>%d</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), seq_num)
    p._p.append(parse_xml(fld_xml))
    
    # Regular desc
    r2 = p.add_run(f" {desc_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = False

def fix_all_50_diagrams_precisely():
    doc = docx.Document('SKRIPSI.docx')
    body = doc._body._body
    
    # Sequence of all 50 BAB IV diagrams and their search markers:
    diagram_specs = [
        # (Marker in heading / text, Caption Description)
        ("Arsitektur sistem usulan menggambarkan", "Arsitektur Sistem Usulan"),
        ("Use Case Diagram digunakan untuk", "Use Case Diagram"),
        ("AD-01", "Activity Diagram Login"),
        ("AD-02", "Activity Diagram Mengelola Data Pengguna"),
        ("AD-03", "Activity Diagram Mengelola Data Guru"),
        ("AD-04", "Activity Diagram Mengelola Kriteria dan Bobot"),
        ("AD-05", "Activity Diagram Mengelola Skala Penilaian"),
        ("AD-06", "Activity Diagram Mengelola Periode Penilaian"),
        ("AD-07", "Activity Diagram Melakukan Penilaian Kinerja Guru"),
        ("AD-08", "Activity Diagram Mencatat Hasil Penilaian"),
        ("AD-09", "Activity Diagram Mengolah Penilaian dengan MOORA"),
        ("AD-10", "Activity Diagram Melihat Pemeringkatan MOORA"),
        ("AD-11", "Activity Diagram Melihat Hasil Penilaian"),
        ("AD-12", "Activity Diagram Melihat Detail Penilaian"),
        ("AD-13", "Activity Diagram Monitoring Kinerja Guru"),
        ("AD-14", "Activity Diagram Melihat Riwayat Penilaian"),
        ("AD-15", "Activity Diagram Mencetak Laporan Penilaian"),
        ("SD-01", "Sequence Diagram Login"),
        ("SD-02", "Sequence Diagram Mengelola Data Pengguna"),
        ("SD-03", "Sequence Diagram Mengelola Data Guru"),
        ("SD-04", "Sequence Diagram Mengelola Kriteria dan Bobot"),
        ("SD-05", "Sequence Diagram Mengelola Skala Penilaian"),
        ("SD-06", "Sequence Diagram Mengelola Periode Penilaian"),
        ("SD-07", "Sequence Diagram Melakukan Penilaian Kinerja Guru"),
        ("SD-08", "Sequence Diagram Mencatat Hasil Penilaian"),
        ("SD-09", "Sequence Diagram Mengolah Penilaian dengan MOORA"),
        ("SD-10", "Sequence Diagram Melihat Pemeringkatan MOORA"),
        ("SD-11", "Sequence Diagram Melihat Hasil Penilaian"),
        ("SD-12", "Sequence Diagram Melihat Detail Penilaian"),
        ("SD-13", "Sequence Diagram Monitoring Kinerja Guru"),
        ("SD-14", "Sequence Diagram Melihat Riwayat Penilaian"),
        ("SD-15", "Sequence Diagram Mencetak Laporan Penilaian"),
        ("Perancangan Class Diagram", "Class Diagram Sistem Pendukung Keputusan MOORA"),
        ("Perancangan Relationship Diagram (ERD)", "Entity Relationship Diagram (ERD)"),
        ("Struktur Menu Administrator", "Struktur Menu Administrator"),
        ("Struktur Menu Kepala Sekolah", "Struktur Menu Kepala Sekolah"),
        ("Struktur Menu Guru", "Struktur Menu Guru"),
        ("Rancangan Antarmuka Halaman Login", "Rancangan Antarmuka Halaman Login"),
        ("Rancangan Antarmuka Dashboard Administrator", "Rancangan Antarmuka Dashboard Administrator"),
        ("Rancangan Antarmuka Kelola Data Guru", "Rancangan Antarmuka Kelola Data Guru"),
        ("Rancangan Antarmuka Kelola Kriteria & Bobot", "Rancangan Antarmuka Kelola Kriteria & Bobot"),
        ("Rancangan Antarmuka Kelola Skala Penilaian / Rubrik", "Rancangan Antarmuka Kelola Skala Penilaian / Rubrik"),
        ("Rancangan Antarmuka Kelola Periode Penilaian", "Rancangan Antarmuka Kelola Periode Penilaian"),
        ("Rancangan Antarmuka Form Transaksi Penilaian Kinerja Guru", "Rancangan Antarmuka Form Transaksi Penilaian Kinerja Guru"),
        ("Rancangan Antarmuka Proses Perhitungan SPK MOORA", "Rancangan Antarmuka Proses Perhitungan SPK MOORA"),
        ("Rancangan Antarmuka Peringkat & Hasil Ranking Guru", "Rancangan Antarmuka Peringkat & Hasil Ranking Guru"),
        ("Rancangan Antarmuka Laporan & Cetak Rekapitulasi", "Rancangan Antarmuka Laporan & Cetak Rekapitulasi"),
        ("Rancangan Antarmuka Kelola Akun Pengguna / User", "Rancangan Antarmuka Kelola Akun Pengguna / User"),
        ("Rancangan Antarmuka Dashboard & Profil Mandiri Guru", "Rancangan Antarmuka Dashboard & Profil Mandiri Guru"),
        ("Flowchart Algoritma MOORA", "Flowchart Proses Komputasi Algoritma Metode MOORA"),
    ]
    
    # 1. Clean up ALL existing "Gambar 4." captions in BAB IV
    for elem in list(body):
        if elem.tag.endswith('p'):
            p_obj = docx.text.paragraph.Paragraph(elem, doc)
            t = p_obj.text.strip()
            xml_s = etree.tostring(elem).decode('utf-8', errors='ignore')
            idx = list(body).index(elem)
            if idx > 780:
                if 'w:drawing' not in xml_s and 'w:pict' not in xml_s:
                    if t.startswith("Gambar 4.") or t.startswith("Gambar 4 ") or "SEQ Gambar" in xml_s:
                        elem.getparent().remove(elem)
                        
    # 2. Iterate through diagram specs and insert caption directly after its drawing
    current_search_idx = 780
    success_count = 0
    
    for seq_num, (marker, desc_text) in enumerate(diagram_specs, start=1):
        found = False
        body_elements = list(body)
        
        # Search for marker paragraph starting from current_search_idx
        for i in range(current_search_idx, len(body_elements)):
            elem = body_elements[i]
            if elem.tag.endswith('p'):
                p_obj = docx.text.paragraph.Paragraph(elem, doc)
                txt = p_obj.text.strip()
                if marker in txt:
                    # Look ahead up to 4 elements for the drawing
                    for off in range(1, 5):
                        if i + off < len(body_elements):
                            cand = body_elements[i + off]
                            cand_xml = etree.tostring(cand).decode('utf-8', errors='ignore')
                            if 'w:drawing' in cand_xml or 'w:pict' in cand_xml:
                                # Insert caption after this drawing!
                                new_p = OxmlElement('w:p')
                                cand.addnext(new_p)
                                new_para = docx.text.paragraph.Paragraph(new_p, doc)
                                set_caption(new_para, seq_num, desc_text)
                                print(f"[OK #{seq_num}] Gambar 4.{seq_num} {desc_text}")
                                current_search_idx = i + off + 1
                                found = True
                                success_count += 1
                                break
                    if found:
                        break
        if not found:
            print(f"[FAILED TO FIND] #{seq_num}: marker='{marker}', desc='{desc_text}'")

    print(f"\nTotal successfully assigned captions: {success_count} / {len(diagram_specs)}")
    doc.save('SKRIPSI.docx')
    print("[SUCCESS] SKRIPSI.docx saved!")

if __name__ == "__main__":
    fix_all_50_diagrams_precisely()
