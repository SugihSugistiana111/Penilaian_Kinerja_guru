import docx
from lxml import etree
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml, OxmlElement
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_caption(p, seq_num, desc_text):
    for child in list(p._p):
        tag = child.tag.split('}')[-1]
        if tag not in ('pPr',):
            p._p.remove(child)
            
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:keepNext')) is None:
        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    
    # Bold prefix
    r1 = p.add_run("Gambar 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    # SEQ field
    fld_xml = '''
        <w:fldSimple %s w:instr=" SEQ Gambar \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="24"/>
                </w:rPr>
                <w:t>%d</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), seq_num)
    p._p.append(parse_xml(fld_xml))
    
    # Regular desc
    r2 = p.add_run(f" {desc_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = False

def fix_all_50_final():
    doc = docx.Document('SKRIPSI.docx')
    body = doc._body._body
    
    diagram_specs = [
        ("Arsitektur sistem usulan menggambarkan", "Arsitektur Sistem Usulan"),
        ("Use Case Diagram digunakan untuk", "Use Case Diagram"),
        ("AD-01\tLogin", "Activity Diagram Login"),
        ("AD-02\tMengelola Data Pengguna", "Activity Diagram Mengelola Data Pengguna"),
        ("AD-03\tMengelola Data Guru", "Activity Diagram Mengelola Data Guru"),
        ("AD-04\tMengelola Kriteria dan Bobot", "Activity Diagram Mengelola Kriteria dan Bobot"),
        ("AD-05\tMengelola Skala Penilaian", "Activity Diagram Mengelola Skala Penilaian"),
        ("AD-06\tMengelola Periode Penilaian", "Activity Diagram Mengelola Periode Penilaian"),
        ("AD-07\tMelakukan Penilaian Kinerja Guru", "Activity Diagram Melakukan Penilaian Kinerja Guru"),
        ("AD-08\tMencatat Hasil Penilaian", "Activity Diagram Mencatat Hasil Penilaian"),
        ("AD-09\tMengolah Penilaian dengan MOORA", "Activity Diagram Mengolah Penilaian dengan MOORA"),
        ("AD-10\tMelihat Pemeringkatan MOORA", "Activity Diagram Melihat Pemeringkatan MOORA"),
        ("AD-11\tMelihat Hasil Penilaian", "Activity Diagram Melihat Hasil Penilaian"),
        ("AD-12\tMelihat Detail Penilaian", "Activity Diagram Melihat Detail Penilaian"),
        ("AD-13\tMonitoring Kinerja Guru", "Activity Diagram Monitoring Kinerja Guru"),
        ("AD-14\tMelihat Riwayat Penilaian", "Activity Diagram Melihat Riwayat Penilaian"),
        ("AD-15\tMencetak Laporan Penilaian", "Activity Diagram Mencetak Laporan Penilaian"),
        ("SD-01\tLogin", "Sequence Diagram Login"),
        ("SD-02\tMengelola Data Pengguna", "Sequence Diagram Mengelola Data Pengguna"),
        ("SD-03\tMengelola Data Guru", "Sequence Diagram Mengelola Data Guru"),
        ("SD-04\tMengelola Kriteria dan Bobot", "Sequence Diagram Mengelola Kriteria dan Bobot"),
        ("SD-05\tMengelola Skala Penilaian", "Sequence Diagram Mengelola Skala Penilaian"),
        ("SD-06\tMengelola Periode Penilaian", "Sequence Diagram Mengelola Periode Penilaian"),
        ("SD-07\tMelakukan Penilaian Kinerja Guru", "Sequence Diagram Melakukan Penilaian Kinerja Guru"),
        ("SD-08\tMencatat Hasil Penilaian", "Sequence Diagram Mencatat Hasil Penilaian"),
        ("SD-09\tMengolah Penilaian dengan MOORA", "Sequence Diagram Mengolah Penilaian dengan MOORA"),
        ("SD-10\tMelihat Pemeringkatan MOORA", "Sequence Diagram Melihat Pemeringkatan MOORA"),
        ("SD-11\tMelihat Hasil Penilaian", "Sequence Diagram Melihat Hasil Penilaian"),
        ("SD-12\tMelihat Detail Penilaian", "Sequence Diagram Melihat Detail Penilaian"),
        ("SD-13\tMonitoring Kinerja Guru", "Sequence Diagram Monitoring Kinerja Guru"),
        ("SD-14\tMelihat Riwayat Penilaian", "Sequence Diagram Melihat Riwayat Penilaian"),
        ("SD-15\tMencetak Laporan Penilaian", "Sequence Diagram Mencetak Laporan Penilaian"),
        ("Perancangan Class Diagram", "Class Diagram Sistem Pendukung Keputusan MOORA"),
        ("Perancangan Relationship Diagram (ERD)", "Entity Relationship Diagram (ERD)"),
        ("Struktur Menu Administrator", "Struktur Menu Administrator"),
        ("Struktur Menu Kepala Sekolah", "Struktur Menu Kepala Sekolah"),
        ("Struktur Menu Guru", "Struktur Menu Guru"),
        ("Rancangan Antarmuka Halaman Login", "Rancangan Antarmuka Halaman Login"),
        ("Rancangan Antarmuka Dashboard Administrator", "Rancangan Antarmuka Dashboard Administrator"),
        ("Rancangan Antarmuka Kelola Data Guru", "Rancangan Antarmuka Kelola Data Guru"),
        ("Rancangan Antarmuka Kelola Kriteria & Bobot", "Rancangan Antarmuka Kelola Kriteria & Bobot"),
        ("Rancangan Antarmuka Kelola Skala Penilaian / Rubrik", "Rancangan Antarmuka Kelola Skala Penilaian / Rubrik"),
        ("Rancangan Antarmuka Kelola Periode Penilaian", "Rancangan Antarmuka Kelola Periode Penilaian"),
        ("Rancangan Antarmuka Form Transaksi Penilaian Kinerja Guru", "Rancangan Antarmuka Form Transaksi Penilaian Kinerja Guru"),
        ("Rancangan Antarmuka Proses Perhitungan SPK MOORA", "Rancangan Antarmuka Proses Perhitungan SPK MOORA"),
        ("Rancangan Antarmuka Peringkat & Hasil Ranking Guru", "Rancangan Antarmuka Peringkat & Hasil Ranking Guru"),
        ("Rancangan Antarmuka Laporan & Cetak Rekapitulasi", "Rancangan Antarmuka Laporan & Cetak Rekapitulasi"),
        ("Rancangan Antarmuka Kelola Akun Pengguna / User", "Rancangan Antarmuka Kelola Akun Pengguna / User"),
        ("Rancangan Antarmuka Dashboard & Profil Mandiri", "Rancangan Antarmuka Dashboard & Profil Mandiri Guru"),
        ("Flowchart Algoritma MOORA", "Flowchart Proses Komputasi Algoritma Metode MOORA"),
    ]
    
    # 1. Clean up ALL existing "Gambar 4." captions in BAB IV
    for elem in list(body):
        if elem.tag.endswith('p'):
            p_obj = docx.text.paragraph.Paragraph(elem, doc)
            t = p_obj.text.strip()
            xml_s = etree.tostring(elem).decode('utf-8', errors='ignore')
            idx = list(body).index(elem)
            if idx > 780:
                if 'w:drawing' not in xml_s and 'w:pict' not in xml_s:
                    if t.startswith("Gambar 4.") or t.startswith("Gambar 4 ") or "SEQ Gambar" in xml_s:
                        elem.getparent().remove(elem)
                        
    # 2. Iterate through diagram specs
    body_elements = list(body)
    success_count = 0
    
    for seq_num, (marker, desc_text) in enumerate(diagram_specs, start=1):
        found = False
        body_elements = list(body)
        
        for i in range(780, len(body_elements)):
            elem = body_elements[i]
            if elem.tag.endswith('p'):
                p_obj = docx.text.paragraph.Paragraph(elem, doc)
                txt = p_obj.text.strip()
                if marker in txt:
                    # Look ahead for drawing
                    for off in range(1, 6):
                        if i + off < len(body_elements):
                            cand = body_elements[i + off]
                            cand_xml = etree.tostring(cand).decode('utf-8', errors='ignore')
                            if 'w:drawing' in cand_xml or 'w:pict' in cand_xml:
                                new_p = OxmlElement('w:p')
                                cand.addnext(new_p)
                                new_para = docx.text.paragraph.Paragraph(new_p, doc)
                                set_caption(new_para, seq_num, desc_text)
                                print(f"[OK #{seq_num}] Gambar 4.{seq_num} {desc_text}")
                                found = True
                                success_count += 1
                                break
                    if found:
                        break
        if not found:
            print(f"[FAILED TO FIND] #{seq_num}: marker='{marker}', desc='{desc_text}'")

    print(f"\nTotal successfully assigned captions: {success_count} / {len(diagram_specs)}")
    doc.save('SKRIPSI.docx')
    print("[SUCCESS] SKRIPSI.docx saved!")

if __name__ == "__main__":
    fix_all_50_final()
