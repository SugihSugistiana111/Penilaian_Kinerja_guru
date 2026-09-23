import docx
from lxml import etree
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml, OxmlElement
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_caption_content(p, label_prefix, seq_type, seq_num, desc_text):
    """
    Format caption:
    - "{label_prefix}" + {seq_num} : BOLD, Times New Roman 12pt
    - " {desc_text}" : REGULAR (NOT BOLD), Times New Roman 12pt
    - Center aligned, space_before=0, space_after=0, keep_with_next=True
    """
    # Clear runs and child elements except pPr
    for child in list(p._p):
        tag = child.tag.split('}')[-1]
        if tag not in ('pPr',):
            p._p.remove(child)
            
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    
    # Ensure pPr has keepNext
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:keepNext')) is None:
        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    
    # Bold prefix: e.g. "Gambar 4."
    r1 = p.add_run(label_prefix)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    # SEQ field with placeholder number
    fld_xml = '''
        <w:fldSimple %s w:instr=" SEQ %s \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="24"/>
                </w:rPr>
                <w:t>%d</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), seq_type, seq_num)
    p._p.append(parse_xml(fld_xml))
    
    # Regular text: e.g. " Rancangan Antarmuka Halaman Login"
    r2 = p.add_run(f" {desc_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = False

def fix_bab4_image_captions():
    doc = docx.Document('SKRIPSI.docx')
    body = doc._body._body
    
    # List of all 50 target images in BAB IV in exact sequential order:
    target_captions = [
        # 1. Arsitektur
        ("Arsitektur", "Arsitektur Sistem Usulan"),
        # 2. Use Case
        ("Use Case", "Use Case Diagram"),
        # 3-17. 15 Activity Diagrams
        ("AD-01", "Activity Diagram Login"),
        ("AD-02", "Activity Diagram Mengelola Data Pengguna"),
        ("AD-03", "Activity Diagram Mengelola Data Guru"),
        ("AD-04", "Activity Diagram Mengelola Kriteria dan Bobot"),
        ("AD-05", "Activity Diagram Mengelola Skala Penilaian"),
        ("AD-06", "Activity Diagram Mengelola Periode Penilaian"),
        ("AD-07", "Activity Diagram Melakukan Penilaian Kinerja Guru"),
        ("AD-08", "Activity Diagram Mencatat Hasil Penilaian"),
        ("AD-09", "Activity Diagram Mengolah Penilaian dengan MOORA"),
        ("AD-10", "Activity Diagram Melihat Pemeringkatan MOORA"),
        ("AD-11", "Activity Diagram Melihat Hasil Penilaian"),
        ("AD-12", "Activity Diagram Melihat Detail Penilaian"),
        ("AD-13", "Activity Diagram Monitoring Kinerja Guru"),
        ("AD-14", "Activity Diagram Melihat Riwayat Penilaian"),
        ("AD-15", "Activity Diagram Mencetak Laporan Penilaian"),
        # 18-32. 15 Sequence Diagrams
        ("SD-01", "Sequence Diagram Login"),
        ("SD-02", "Sequence Diagram Mengelola Data Pengguna"),
        ("SD-03", "Sequence Diagram Mengelola Data Guru"),
        ("SD-04", "Sequence Diagram Mengelola Kriteria dan Bobot"),
        ("SD-05", "Sequence Diagram Mengelola Skala Penilaian"),
        ("SD-06", "Sequence Diagram Mengelola Periode Penilaian"),
        ("SD-07", "Sequence Diagram Melakukan Penilaian Kinerja Guru"),
        ("SD-08", "Sequence Diagram Mencatat Hasil Penilaian"),
        ("SD-09", "Sequence Diagram Mengolah Penilaian dengan MOORA"),
        ("SD-10", "Sequence Diagram Melihat Pemeringkatan MOORA"),
        ("SD-11", "Sequence Diagram Melihat Hasil Penilaian"),
        ("SD-12", "Sequence Diagram Melihat Detail Penilaian"),
        ("SD-13", "Sequence Diagram Monitoring Kinerja Guru"),
        ("SD-14", "Sequence Diagram Melihat Riwayat Penilaian"),
        ("SD-15", "Sequence Diagram Mencetak Laporan Penilaian"),
        # 33. Class Diagram
        ("Class Diagram", "Class Diagram Sistem Pendukung Keputusan MOORA"),
        # 34. ERD
        ("ERD", "Entity Relationship Diagram (ERD)"),
        # 35-37. Struktur Menu
        ("Menu Admin", "Struktur Menu Administrator"),
        ("Menu Kepsek", "Struktur Menu Kepala Sekolah"),
        ("Menu Guru", "Struktur Menu Guru"),
        # 38-49. 12 Rancangan Antarmuka
        ("UI-01", "Rancangan Antarmuka Halaman Login"),
        ("UI-02", "Rancangan Antarmuka Dashboard Administrator"),
        ("UI-03", "Rancangan Antarmuka Kelola Data Guru"),
        ("UI-04", "Rancangan Antarmuka Kelola Kriteria & Bobot"),
        ("UI-05", "Rancangan Antarmuka Kelola Skala Penilaian / Rubrik"),
        ("UI-06", "Rancangan Antarmuka Kelola Periode Penilaian"),
        ("UI-07", "Rancangan Antarmuka Form Transaksi Penilaian Kinerja Guru"),
        ("UI-08", "Rancangan Antarmuka Proses Perhitungan SPK MOORA"),
        ("UI-09", "Rancangan Antarmuka Peringkat & Hasil Ranking Guru"),
        ("UI-10", "Rancangan Antarmuka Laporan & Cetak Rekapitulasi"),
        ("UI-11", "Rancangan Antarmuka Kelola Akun Pengguna / User"),
        ("UI-12", "Rancangan Antarmuka Dashboard & Profil Mandiri Guru"),
        # 50. Flowchart MOORA
        ("Flowchart", "Flowchart Proses Komputasi Algoritma Metode MOORA"),
    ]
    
    print(f"Total target BAB IV image captions to assign: {len(target_captions)}")
    
    # Step 1: Clean up ALL existing "Gambar 4." captions in BAB IV
    # so we can rebuild them cleanly directly underneath each drawing!
    
    body_elements = list(body)
    cleaned_old = 0
    for i, elem in enumerate(body_elements):
        tag = elem.tag.split('}')[-1]
        if tag == 'p':
            p_obj = docx.text.paragraph.Paragraph(elem, doc)
            t = p_obj.text.strip()
            # Only remove if it's in BAB IV and starts with "Gambar 4."
            if (t.startswith("Gambar 4.") or t.startswith("Gambar 4 ")) and i > 750:
                # Check if it has drawing - don't remove if paragraph itself contains drawing!
                if 'w:drawing' not in elem.xml and 'w:pict' not in elem.xml:
                    elem.getparent().remove(elem)
                    cleaned_old += 1
    
    print(f"Cleaned {cleaned_old} old/inconsistent Gambar 4.x captions in BAB IV")
    
    # Step 2: Now find every drawing in BAB IV and attach the proper caption directly underneath it!
    body_elements = list(body)
    
    # Find drawing paragraphs from BAB IV onwards (index > 750)
    bab4_drawings = []
    for i, elem in enumerate(body_elements):
        if i > 750:
            xml_str = etree.tostring(elem).decode('utf-8', errors='ignore')
            if 'w:drawing' in xml_str or 'w:pict' in xml_str:
                bab4_drawings.append(i)
                
    print(f"Found {len(bab4_drawings)} drawings in BAB IV:")
    
    # Check match
    if len(bab4_drawings) != len(target_captions):
        print(f"[WARNING] Drawings count ({len(bab4_drawings)}) differs from targets ({len(target_captions)})")
        # Let's inspect context of each drawing to verify mapping
        for idx, pos in enumerate(bab4_drawings):
            e = body_elements[pos]
            # Context
            ctx = ""
            for off in range(-2, 2):
                if 0 <= pos + off < len(body_elements):
                    ee = body_elements[pos + off]
                    if ee.tag.endswith('p'):
                        po = docx.text.paragraph.Paragraph(ee, doc)
                        if po.text.strip():
                            ctx += f" [{po.text.strip()[:35]}]"
            print(f"  Drawing #{idx+1} at pos {pos}: {ctx}")
    
    # Step 3: Insert the caption directly after each drawing element
    for seq_num, (key, desc_text) in enumerate(target_captions, start=1):
        if seq_num - 1 < len(bab4_drawings):
            drawing_elem_idx = bab4_drawings[seq_num - 1]
            drawing_elem = body_elements[drawing_elem_idx]
            
            # Create new caption paragraph directly AFTER the drawing element
            new_p = OxmlElement('w:p')
            drawing_elem.addnext(new_p)
            
            new_para = docx.text.paragraph.Paragraph(new_p, doc)
            set_caption_content(new_para, "Gambar 4.", "Gambar", seq_num, desc_text)
            print(f"[OK] Gambar 4.{seq_num} {desc_text}")
    
    # Step 4: Save
    try:
        doc.save('SKRIPSI.docx')
        print("[SUCCESS] SKRIPSI.docx berhasil disimpan dengan 50 caption gambar berurutan!")
    except PermissionError:
        doc.save('SKRIPSI_Captions_Fixed.docx')
        print("[WARNING] SKRIPSI.docx sedang terbuka. Disimpan sebagai SKRIPSI_Captions_Fixed.docx")

if __name__ == "__main__":
    fix_bab4_image_captions()
