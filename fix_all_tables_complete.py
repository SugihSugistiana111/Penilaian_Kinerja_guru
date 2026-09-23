import docx
from lxml import etree
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml, OxmlElement
from docx.shared import Pt

def fix_all_tables():
    doc = docx.Document('SKRIPSI.docx')
    body_elements = list(doc._body._body)
    
    fixed_floating = 0
    fixed_headers = 0
    fixed_cant_split = 0
    fixed_captions = 0
    
    # 1. Fix all tables in the document
    for tbl in doc.tables:
        tbl_elem = tbl._tbl
        tblPr = tbl_elem.find(qn('w:tblPr'))
        
        # Remove floating positioning (tblpPr) so table is inline and repeat header works
        if tblPr is not None:
            tblpPr = tblPr.find(qn('w:tblpPr'))
            if tblpPr is not None:
                tblPr.remove(tblpPr)
                fixed_floating += 1
        
        # Ensure row 0 has tblHeader
        if len(tbl.rows) > 0:
            tr0 = tbl.rows[0]._tr
            trPr0 = tr0.get_or_add_trPr()
            tblHeader = trPr0.find(qn('w:tblHeader'))
            if tblHeader is None:
                trPr0.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
                fixed_headers += 1
        
        # Ensure all rows have cantSplit
        for row in tbl.rows:
            trPr = row._tr.get_or_add_trPr()
            cantSplit = trPr.find(qn('w:cantSplit'))
            if cantSplit is None:
                trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
                fixed_cant_split += 1

    # 2. Fix all table and image captions
    for p in doc.paragraphs:
        xml = p._p.xml
        if 'SEQ Tabel' in xml or 'SEQ Gambar' in xml or p.text.strip().startswith('Tabel 4.') or p.text.strip().startswith('Gambar 4.'):
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.keep_with_next = True
            
            # Ensure pPr has keepNext
            pPr = p._p.get_or_add_pPr()
            kwn = pPr.find(qn('w:keepNext'))
            if kwn is None:
                pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
            fixed_captions += 1
            
    print(f"[RESULTS]")
    print(f"- Fixed floating tables (removed tblpPr): {fixed_floating}")
    print(f"- Added tblHeader to tables: {fixed_headers}")
    print(f"- Added cantSplit to rows: {fixed_cant_split}")
    print(f"- Updated captions (spacing=0, keepNext=True): {fixed_captions}")
    
    try:
        doc.save('SKRIPSI.docx')
        print("[SUCCESS] SKRIPSI.docx berhasil disimpan!")
    except PermissionError:
        doc.save('SKRIPSI_TableFix.docx')
        print("[WARNING] SKRIPSI.docx sedang terbuka. Disimpan ke SKRIPSI_TableFix.docx")

if __name__ == "__main__":
    fix_all_tables()
