import docx
import re
from lxml import etree
from docx.oxml.ns import nsdecls, qn
from docx.oxml import parse_xml
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_table_caption(p, seq_num, desc_text):
    for child in list(p._p):
        tag = child.tag.split('}')[-1]
        if tag not in ('pPr',):
            p._p.remove(child)
            
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.keep_with_next = True
    
    pPr = p._p.get_or_add_pPr()
    if pPr.find(qn('w:keepNext')) is None:
        pPr.append(parse_xml(f'<w:keepNext {nsdecls("w")}/>'))
    
    # Bold prefix: "Tabel 4."
    r1 = p.add_run("Tabel 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    # SEQ field (BOLD)
    fld_xml = '''
        <w:fldSimple %s w:instr=" SEQ Tabel \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="24"/>
                </w:rPr>
                <w:t>%d</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), seq_num)
    p._p.append(parse_xml(fld_xml))
    
    # Regular desc (NOT BOLD)
    r2 = p.add_run(f" {desc_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = False

def fix_all_tables_exact():
    doc = docx.Document('SKRIPSI.docx')
    
    target_tables = [
        (r'Kebutuhan\s+Perangkat\s+Keras', "Kebutuhan Perangkat Keras"),
        (r'Kebutuhan\s+Perangkat\s+Lunak', "Kebutuhan Perangkat Lunak"),
        (r'Identifikasi\s+Aktor', "Identifikasi Aktor"),
        (r'Pemetaan\s+Kebutuhan\s+Fungsional', "Pemetaan Kebutuhan Fungsional terhadap Use Case"),
        (r'Deskripsi\s+Use\s+Case', "Deskripsi Use Case"),
        (r'Skenario\s+Use\s+Case\s+Login', "Skenario Use Case Login"),
        (r'Skenario\s+Use\s+Case\s+Mengelola\s+Data\s+Pengguna', "Skenario Use Case Mengelola Data Pengguna"),
        (r'Skenario\s+Use\s+Case\s+Mengelola\s+Data\s+Guru', "Skenario Use Case Mengelola Data Guru"),
        (r'Skenario\s+Use\s+Case\s+Mengelola\s+Kriteria', "Skenario Use Case Mengelola Kriteria dan Bobot"),
        (r'Skenario\s+Use\s+Case\s+Mengelola\s+Skala', "Skenario Use Case Mengelola Skala Penilaian"),
        (r'Skenario\s+Use\s+Case\s+Mengelola\s+Periode', "Skenario Use Case Mengelola Periode Penilaian"),
        (r'Skenario\s+Use\s+Case\s+Melakukan\s+Penilaian', "Skenario Use Case Melakukan Penilaian Kinerja Guru"),
        (r'Skenario\s+Use\s+Case\s+Mencatat\s+Hasil\s+Penilaian', "Skenario Use Case Mencatat Hasil Penilaian"),
        (r'Skenario\s+Use\s+Case\s+Mengolah\s+Penilaian', "Skenario Use Case Mengolah Penilaian dengan MOORA"),
        (r'Skenario\s+Use\s+Case\s+Melihat\s+Hasil\s+Penilaian', "Skenario Use Case Melihat Hasil Penilaian"),
        (r'Skenario\s+Use\s+Case\s+Melihat\s+Pemeringkatan', "Skenario Use Case Melihat Pemeringkatan"),
        (r'Skenario\s+Use\s+Case\s+Monitoring\s+Kinerja', "Skenario Use Case Monitoring Kinerja Guru"),
        (r'Skenario\s+Use\s+Case\s+Melihat\s+Detail\s+Penilaian', "Skenario Use Case Melihat Detail Penilaian"),
        (r'Skenario\s+Use\s+Case\s+Melihat\s+Riwayat\s+Penilaian', "Skenario Use Case Melihat Riwayat Penilaian"),
        (r'Skenario\s+Use\s+Case\s+Mencetak\s+Laporan', "Skenario Use Case Mencetak Laporan Penilaian"),
        (r'Struktur\s+Tabel\s+Role', "Struktur Tabel Role (tbl_role)"),
        (r'Struktur\s+Tabel\s+User', "Struktur Tabel User (tbl_user)"),
        (r'Struktur\s+Tabel\s+Guru', "Struktur Tabel Guru (tbl_guru)"),
        (r'Struktur\s+Tabel\s+Kriteria', "Struktur Tabel Kriteria (tbl_kriteria)"),
        (r'Struktur\s+Tabel\s+Skala', "Struktur Tabel Skala Penilaian (tbl_skala_penilaian)"),
        (r'Struktur\s+Tabel\s+Periode', "Struktur Tabel Periode Penilaian (tbl_periode_penilaian)"),
        (r'Struktur\s+Tabel\s+Penilaian\b', "Struktur Tabel Penilaian (tbl_penilaian)"),
        (r'Struktur\s+Tabel\s+Detail\s+Penilaian', "Struktur Tabel Detail Penilaian (tbl_detail_penilaian)"),
        (r'Struktur\s+Tabel\s+Hasil\s+MOORA', "Struktur Tabel Hasil MOORA (tbl_hasil_moora)"),
        (r'Struktur\s+Tabel\s+Detail\s+MOORA', "Struktur Tabel Detail MOORA (tbl_detail_moora)"),
        (r'Struktur\s+Tabel\s+Log', "Struktur Tabel Log Aktivitas (tbl_log_aktivitas)"),
        (r'Daftar\s+Kriteria,\s+Bobot', "Daftar Kriteria, Bobot Preferensi, dan Sifat Atribut MOORA"),
        (r'Daftar\s+Alternatif\s+Guru', "Daftar Alternatif Guru SMA Al-Ihsan Boarding School"),
        (r'Matriks\s+Keputusan\s+\(X\)', "Matriks Keputusan (X) Evaluasi Kinerja Guru"),
        (r'Matriks\s+Ternormalisasi\s+\(X\*\)', "Matriks Ternormalisasi (X*) Metode MOORA"),
        (r'Matriks\s+Normalisasi\s+Terbobot', "Matriks Normalisasi Terbobot (W x X*)"),
        (r'Hasil\s+Akhir\s+Nilai\s+Preferensi', "Hasil Akhir Nilai Preferensi (Yi) dan Peringkat Kinerja Guru (MOORA)"),
    ]
    
    # Process BAB IV paragraphs (index > 740)
    success = 0
    for seq_num, (pattern, desc_clean) in enumerate(target_tables, start=1):
        found = False
        for i, p in enumerate(doc.paragraphs):
            if i > 740:
                t = p.text.strip()
                if re.search(pattern, t, re.IGNORECASE):
                    set_table_caption(p, seq_num, desc_clean)
                    print(f"[OK #{seq_num:2d}] Tabel 4.{seq_num:2d} {desc_clean}")
                    found = True
                    success += 1
                    break
        if not found:
            print(f"[FAILED #{seq_num:2d}] pattern='{pattern}'")
            
    print(f"\nTotal exact table captions matched: {success} / {len(target_tables)}")
    doc.save('SKRIPSI.docx')
    print("[SUCCESS] SKRIPSI.docx saved!")

if __name__ == "__main__":
    fix_all_tables_exact()
