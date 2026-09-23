import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_image_caption(paragraph, num_suffix, title_text, label="Gambar"):
    """
    Menambahkan caption Word otomatis dengan SEQ field yang kompatibel 
    dengan References -> Insert Table of Figures (Daftar Gambar).
    """
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.keep_with_next = True
    
    r1 = paragraph.add_run(f"Gambar 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    fld_xml = r'''
        <w:fldSimple %s w:instr="SEQ Gambar \* ARABIC">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:t>%s</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), num_suffix)
    paragraph._p.append(parse_xml(fld_xml))
    
    r2 = paragraph.add_run(f" {title_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.bold = True

def generate_antarmuka_document():
    doc = docx.Document()
    
    # Page setup - A4, Margin Skripsi (Top 3cm, Left 4cm, Bottom 3cm, Right 3cm)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(3.0)
        section.left_margin = Cm(4.0)
        section.right_margin = Cm(3.0)
        
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)
    
    # Header Bab IV
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_h1 = h1.add_run("BAB IV\nHASIL DAN PEMBAHASAN\n4.4 PERANCANGAN ANTARMUKA PENGGUNA (USER INTERFACE DESIGN)")
    r_h1.font.name = 'Times New Roman'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    h1.paragraph_format.space_after = Pt(18)
    
    # Intro
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_intro = p_intro.add_run(
        "Perancangan antarmuka pengguna (User Interface Design) pada Sistem Pendukung Keputusan Penilaian Kinerja Guru "
        "SMA Al-Ihsan Boarding School disusun dalam bentuk rancangan wireframe/mockup monokromatik (tahap perancangan) "
        "yang merepresentasikan tata letak komponen, hierarki informasi, dan alur kontrol interaksi pengguna secara identik "
        "dengan implementasi sistem. Perancangan antarmuka dikelompokkan berdasarkan modul hak akses (Administrator, "
        "Kepala Sekolah, dan Guru) sebagai berikut:"
    )
    r_intro.font.name = 'Times New Roman'
    r_intro.font.size = Pt(12)
    
    # Data 12 interfaces
    ui_data = [
        {
            "num": "1",
            "title": "Rancangan Antarmuka Halaman Login",
            "img": "wireframes/wf_01_login.png",
            "fungsi": "Halaman Login berfungsi sebagai gerbang autentikasi dan otorisasi terpusat (multilevel login) bagi seluruh aktor pengguna sistem (Administrator, Kepala Sekolah, dan Guru) untuk memvalidasi kredensial username dan password sebelum mengakses hak fitur masing-masing.",
            "hak_akses": "Administrator, Kepala Sekolah, Guru",
            "controls": [
                ["1", "Logo Al-Ihsan", "Image / Brand", "Menampilkan identitas resmi SMA Al-Ihsan Boarding School"],
                ["2", "Field Username", "Input Text", "Menerima input username unik akun pengguna terdaftar"],
                ["3", "Field Password", "Input Password", "Menerima input password terenkripsi pengguna"],
                ["4", "Tombol Masuk ke Sistem", "Button", "Mengeksekusi proses verifikasi kredensial autentikasi login"]
            ]
        },
        {
            "num": "2",
            "title": "Rancangan Antarmuka Dashboard Administrator",
            "img": "wireframes/wf_02_dashboard_admin.png",
            "fungsi": "Halaman Dashboard Administrator berfungsi sebagai pusat kendali utama yang menampilkan ringkasan data statistik sistem secara real-time, status progres kelengkapan penilaian kriteria guru pada periode berjalan, peringkat 3 besar guru terbaik, serta rekaman log aktivitas terbaru.",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Sidebar Navigasi", "Navigation Menu", "Menu peralihan antar modul data master, penilaian, SPK MOORA, dan laporan"],
                ["2", "4 Kartu Statistik", "Info Card Widget", "Menampilkan total Guru Aktif, Sudah Dinilai, Belum Dinilai, dan Periode Aktif"],
                ["3", "Status Kelengkapan Penilaian", "Progress Bar & Alert", "Menampilkan persentase kesiapan data nilai guru untuk kalkulasi MOORA"],
                ["4", "Tombol Hitung MOORA", "Action Button", "Akses pintas langsung untuk menjalankan kalkulasi optimasi MOORA"],
                ["5", "Top 3 Leaderboard Guru", "List Widget", "Menampilkan 3 guru dengan nilai preferensi Yi tertinggi"],
                ["6", "Log Aktivitas Terbaru", "Audit Trail List", "Menampilkan 5 riwayat aktivitas manipulasi data sistem terbaru"]
            ]
        },
        {
            "num": "3",
            "title": "Rancangan Antarmuka Kelola Data Guru",
            "img": "wireframes/wf_03_kelola_guru.png",
            "fungsi": "Halaman Kelola Data Guru berfungsi untuk mengelola seluruh data master guru SMA Al-Ihsan Boarding School yang menjadi objek alternatif penilaian kinerja, meliputi identitas NIP, jenis kelamin, tugas mengajar, jumlah jam tatap muka, dan rombongan belajar.",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Pencarian & Filter Status", "Input Search & Dropdown", "Menyaring daftar guru berdasarkan kata kunci nama/NIP dan status aktif"],
                ["2", "Tombol + Tambah Guru", "Primary Button", "Membuka modal dialog form pendaftaran data master guru baru"],
                ["3", "Tabel Data Guru", "Data Table Grid", "Menampilkan kolom No, NIP, Nama, JK, Tugas Mengajar, Jam Ajar, Rombel, Status"],
                ["4", "Tombol Aksi [Edit] & [Hapus]", "Action Links", "Melakukan pengeditan data guru atau menghapus data alternatif"],
                ["5", "Modal Form Input Guru", "Modal Dialog Form", "Form isian NIP, nama lengkap, penugasan mengajar, jam ajar, dan rombel"]
            ]
        },
        {
            "num": "4",
            "title": "Rancangan Antarmuka Kelola Kriteria & Bobot",
            "img": "wireframes/wf_04_kelola_kriteria.png",
            "fungsi": "Halaman Kelola Kriteria berfungsi untuk mengatur data kriteria evaluasi kinerja guru (C1-Pedagogik, C2-Profesional, C3-Kepribadian, C4-Sosial), menetapkan nilai bobot preferensi (total 1.00 / 100%), dan menentukan jenis sifat kriteria metode MOORA (Benefit atau Cost).",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Banner Status Total Bobot", "Alert Info Box", "Memvalidasi apakah akumulasi seluruh bobot kriteria telah genap bernilai 1.00"],
                ["2", "Tabel Kriteria Penilaian", "Data Table Grid", "Menampilkan Kode (C1-C4), Nama Kriteria, Nilai Bobot, Persentase, Sifat MOORA"],
                ["3", "Tombol Aksi [Edit / Atur]", "Action Button", "Membuka form penyesuaian bobot preferensi kriteria"],
                ["4", "Form Modal Bobot Kriteria", "Form Input Container", "Form pengaturan kode kriteria, nama, nilai bobot numerik, dan jenis Benefit/Cost"]
            ]
        },
        {
            "num": "5",
            "title": "Rancangan Antarmuka Kelola Skala Penilaian / Rubrik",
            "img": "wireframes/wf_05_skala_penilaian.png",
            "fungsi": "Halaman Kelola Skala Penilaian berfungsi untuk mengelola rubrik sub-kriteria dan pedoman indikator skor kuantitatif (skala 1 sampai 5) dari masing-masing kriteria penilaian agar proses evaluasi guru berjalan objektif dan terstandarisasi.",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Dropdown Pilihan Kriteria", "Select Dropdown", "Memilih kriteria spesifik yang akan dikonfigurasi skala rubriknya"],
                ["2", "Tabel Rubrik Penilaian", "Data Table Grid", "Menampilkan tingkatan Nilai (1-5), Label Predikat, dan Deskripsi Indikator Capaian"],
                ["3", "Tombol [Edit] & [Hapus]", "Action Links", "Memodifikasi uraian kalimat indikator ketercapaian rubrik kompetensi"]
            ]
        },
        {
            "num": "6",
            "title": "Rancangan Antarmuka Kelola Periode Penilaian",
            "img": "wireframes/wf_06_periode_penilaian.png",
            "fungsi": "Halaman Kelola Periode Penilaian berfungsi untuk membuka dan menutup gelombang evaluasi kinerja guru secara berkala (bulanan/semesteran), mengatur periode berstatus AKTIF, serta mengarsipkan riwayat periode yang telah SELESAI.",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Tombol + Tambah Periode", "Primary Button", "Membuka form pembuatan periode evaluasi baru (Bulan, Tahun, Nama Periode)"],
                ["2", "Tabel Periode Penilaian", "Data Table Grid", "Menampilkan Nama Periode, Tahun, Status (DRAFT/AKTIF/SELESAI), Progres Dinilai"],
                ["3", "Tombol [Tutup / Buka Periode]", "Status Toggle Button", "Mengubah status operasional periode penilaian agar data terkunci aman"]
            ]
        },
        {
            "num": "7",
            "title": "Rancangan Antarmuka Form Transaksi Penilaian Kinerja Guru",
            "img": "wireframes/wf_07_form_penilaian.png",
            "fungsi": "Halaman Form Penilaian berfungsi sebagai antarmuka pengisian skor evaluasi kinerja setiap guru pada periode aktif, di mana evaluator menginput nilai kuantitatif untuk kriteria Pedagogik, Profesional, Kepribadian, dan Sosial disertai catatan rekomendasi.",
            "hak_akses": "Kepala Sekolah, Administrator",
            "controls": [
                ["1", "Informasi Profil Guru", "Header Profile Card", "Menampilkan NIP, Nama Guru, Mata Pelajaran, Beban Jam Ajar, dan Nama Evaluator"],
                ["2", "Kartu Input Skor Kriteria (C1-C4)", "Radio Group / Scale", "Pilihan opsi skala nilai (1-5) dengan label keterangan ketercapaian"],
                ["3", "Field Catatan Evaluator", "Textarea Multi-line", "Menerima input catatan kualitatif, saran pengembangan, atau apresiasi penilai"],
                ["4", "Tombol [Simpan Draft]", "Secondary Button", "Menyimpan nilai sementara tanpa memvalidasi kelengkapan final"],
                ["5", "Tombol [Simpan Selesai]", "Primary Button", "Menyimpan nilai evaluasi final dan memperbarui status menjadi SELESAI"]
            ]
        },
        {
            "num": "8",
            "title": "Rancangan Antarmuka Proses Perhitungan SPK MOORA",
            "img": "wireframes/wf_08_proses_moora.png",
            "fungsi": "Halaman Perhitungan MOORA berfungsi untuk mengeksekusi algoritma SPK MOORA secara transparan melalui tahapan pembentukan Matriks Keputusan (X), Matriks Ternormalisasi (X*), Matriks Terbobot (W*X*), hingga kalkulasi Nilai Preferensi Optimasi Multiobjektif (Yi).",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Status Validasi Data Guru", "Status Badge Box", "Menampilkan indikator kesiapan kelengkapan data nilai seluruh guru aktif"],
                ["2", "Tombol [PROSES KALKULASI]", "Primary Action Button", "Menjalankan kalkulasi matematis SPK metode MOORA secara otomatis"],
                ["3", "Tab Selektor Matriks", "Tab Navigation", "Pilihan tampilan: 1. Matriks X, 2. Matriks X*, 3. Matriks W*X*, 4. Nilai Yi"],
                ["4", "Tabel Matriks MOORA", "Data Table Grid", "Menampilkan nilai kalkulasi numerik per kriteria untuk setiap guru alternatif"],
                ["5", "Ringkasan Hasil Optimasi", "Summary Card", "Menampilkan penjelasan nilai preferensi Yi dan formula max-min MOORA"]
            ]
        },
        {
            "num": "9",
            "title": "Rancangan Antarmuka Peringkat & Hasil Ranking Guru",
            "img": "wireframes/wf_09_ranking_guru.png",
            "fungsi": "Halaman Ranking Guru berfungsi untuk menyajikan hasil akhir perankingan kelayakan kinerja guru yang telah diurutkan dari nilai preferensi Yi tertinggi hingga terendah, dilengkapi visualisasi podium 3 besar guru terbaik.",
            "hak_akses": "Administrator, Kepala Sekolah",
            "controls": [
                ["1", "Podium Top 3 Terbaik", "Podium Display Widget", "Menampilkan kartu visual Juara 1, Juara 2, dan Juara 3 beserta skor Yi"],
                ["2", "Tabel Peringkat Lengkap", "Data Table Grid", "Menampilkan kolom Peringkat (Rank 1..N), NIP, Nama, Mapel, Nilai Yi, Predikat"],
                ["3", "Badge Predikat Kinerja", "Status Badge", "Label kategori capaian kinerja (Sangat Baik, Baik, Cukup, Kurang)"]
            ]
        },
        {
            "num": "10",
            "title": "Rancangan Antarmuka Laporan & Cetak Rekapitulasi",
            "img": "wireframes/wf_10_laporan_cetak.png",
            "fungsi": "Halaman Laporan berfungsi untuk mencetak dan mengekspor rekapitulasi komprehensif hasil evaluasi penilaian kinerja guru dan perankingan MOORA ke dalam format dokumen resmi PDF (siap tanda tangan) dan file spreadsheet Microsoft Excel.",
            "hak_akses": "Administrator, Kepala Sekolah",
            "controls": [
                ["1", "Filter Periode & Kategori", "Dropdown Filter", "Menyaring laporan berdasarkan periode evaluasi dan kategori guru"],
                ["2", "Tombol [Cetak Laporan PDF]", "Action Button", "Menghasilkan format cetak dokumen resmi PDF lengkap dengan lembar pengesahan"],
                ["3", "Tombol [Export Excel]", "Action Button", "Mengunduh rekapitulasi data penilaian dalam format spreadsheet (.xlsx)"],
                ["4", "Lembar Pratinjau Dokumen", "Document Preview Box", "Pratinjau tampilan tabel rekapitulasi nilai C1-C4, ranking, dan kolom tanda tangan"]
            ]
        },
        {
            "num": "11",
            "title": "Rancangan Antarmuka Kelola Akun Pengguna / User",
            "img": "wireframes/wf_11_kelola_pengguna.png",
            "fungsi": "Halaman Kelola Pengguna berfungsi untuk mengatur data akun login sistem, menetapkan tingkat hak akses (Role ADMIN, KEPALA_SEKOLAH, atau GURU), menghubungkan akun guru dengan data master guru, serta melakukan reset password.",
            "hak_akses": "Administrator",
            "controls": [
                ["1", "Tombol + Tambah User", "Primary Button", "Membuka form pembuatan akun kredensial pengguna baru"],
                ["2", "Tabel Akun Pengguna", "Data Table Grid", "Menampilkan Nama Pengguna, Username, Hak Akses (Role), Relasi Guru, Status"],
                ["3", "Tombol [Reset Password]", "Action Link", "Mereset password pengguna ke kata sandi standar secara instan"],
                ["4", "Tombol [Edit Akun]", "Action Link", "Memodifikasi hak akses atau menonaktifkan akun pengguna"]
            ]
        },
        {
            "num": "12",
            "title": "Rancangan Antarmuka Dashboard & Profil Mandiri Guru",
            "img": "wireframes/wf_12_dashboard_guru.png",
            "fungsi": "Halaman Dashboard Guru berfungsi sebagai portal mandiri bagi tenaga pendidik (guru) untuk melihat biodata profil dinas, beban jam tatap muka, rincian skor capaian kriteria evaluasi (C1-C4), peringkat kinerja mandiri, serta catatan evaluasi dari Kepala Sekolah.",
            "hak_akses": "Guru (Tenaga Pendidik)",
            "controls": [
                ["1", "Kartu Profil & Beban Kerja", "Biodata Card", "Menampilkan NIP, Nama, Mapel, Tugas Tambahan, Jam Ajar, dan Kelas yang Diampu"],
                ["2", "4 Widget Capaian Kriteria", "Score Badge Cards", "Menampilkan nilai skor perolehan C1, C2, C3, dan C4 beserta predikat capaian"],
                ["3", "Kotak Hasil Akhir MOORA", "Result Summary Card", "Menampilkan nilai preferensi Yi yang diperoleh dan posisi peringkat guru"],
                ["4", "Catatan Evaluasi Kepala Sekolah", "Feedback Box", "Menampilkan pesan umpan balik kualitatif dan saran peningkatan kinerja"]
            ]
        }
    ]

    col_widths = [Cm(0.9), Cm(4.8), Cm(3.2), Cm(6.1)]
    headers = ["No", "Nama Kontrol / Komponen", "Tipe Kontrol", "Fungsi / Keterangan"]

    for idx, item in enumerate(ui_data):
        # Heading 4.4.x
        p_sec = doc.add_paragraph()
        p_sec.paragraph_format.line_spacing = 1.5
        p_sec.paragraph_format.space_before = Pt(14)
        p_sec.paragraph_format.space_after = Pt(4)
        p_sec.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        r_sec = p_sec.add_run(f"4.4.{item['num']} {item['title']}\n")
        r_sec.font.name = 'Times New Roman'
        r_sec.font.size = Pt(12)
        r_sec.font.bold = True
        
        # Narrative Deskripsi & Hak Akses
        r_desc = p_sec.add_run(f"{item['fungsi']} Antarmuka ini dapat diakses oleh pengguna dengan hak akses ")
        r_role = p_sec.add_run(f"{item['hak_akses']}.")
        r_role.font.bold = True
        p_sec.add_run(" Rancangan tampilan antarmuka disajikan pada gambar berikut:")

        # Embed Wireframe Image
        if os.path.exists(item['img']):
            p_img = doc.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(8)
            p_img.paragraph_format.space_after = Pt(4)
            p_img.paragraph_format.keep_with_next = True
            
            run_img = p_img.add_run()
            run_img.add_picture(item['img'], width=Cm(15.2))
            
            # Word Caption
            p_cap = doc.add_paragraph(style='Caption')
            add_image_caption(p_cap, item['num'], item['title'])
            
        # Paragraf Pengantar Tabel Kontrol
        p_tbl_intro = doc.add_paragraph()
        p_tbl_intro.paragraph_format.line_spacing = 1.5
        p_tbl_intro.paragraph_format.space_before = Pt(6)
        p_tbl_intro.paragraph_format.space_after = Pt(4)
        p_tbl_intro.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r_ti = p_tbl_intro.add_run(
            f"Fungsi dan kegunaan dari setiap elemen kendali/kontrol yang terdapat pada {item['title'].lower()} "
            f"dijelaskan secara rinci pada tabel di bawah ini:"
        )
        r_ti.font.name = 'Times New Roman'

        # Component Explanation Table
        num_rows = len(item['controls']) + 1
        table = doc.add_table(rows=num_rows, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, color="B0B0B0", sz="4", val="single")

        # Header Repeat & CantSplit
        header_tr = table.rows[0]._tr.get_or_add_trPr()
        header_tr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
        header_tr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

        # Header Content
        for col_idx, h_text in enumerate(headers):
            cell = table.cell(0, col_idx)
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, "EAECEF")
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            
            run = p.add_run(h_text)
            run.font.name = 'Times New Roman'
            run.font.size = Pt(10.5)
            run.font.bold = True

        # Data Rows
        for r_idx, row_data in enumerate(item['controls']):
            row_elem = table.rows[r_idx + 1]
            trPr = row_elem._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

            bg_color = "FAFAFA" if r_idx % 2 == 1 else "FFFFFF"

            for c_idx, val in enumerate(row_data):
                cell = row_elem.cells[c_idx]
                cell.width = col_widths[c_idx]
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_background(cell, bg_color)
                set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
                
                p = cell.paragraphs[0]
                p.paragraph_format.line_spacing = 1.15
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                
                if c_idx == 0:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                elif c_idx == 2:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                
                run = p.add_run(val)
                run.font.name = 'Times New Roman'
                run.font.size = Pt(10)
                if c_idx == 1:
                    run.font.bold = True

        # Spacing after each section
        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(4)
        p_space.paragraph_format.space_after = Pt(8)

    # Save to DOCX
    out_docx = r"c:\Penilaian_Kinerja_Guru\Bab4_Perancangan_Antarmuka.docx"
    try:
        doc.save(out_docx)
        print(f"[SUCCESS] File DOCX berhasil disimpan di: {out_docx}")
    except PermissionError:
        alt_docx = r"c:\Penilaian_Kinerja_Guru\Bab4_Perancangan_Antarmuka_Baru.docx"
        doc.save(alt_docx)
        print(f"[WARNING] File Bab4_Perancangan_Antarmuka.docx sedang dibuka di Word.")
        print(f"[SUCCESS] Disimpan sebagai: {alt_docx}")

    # Generate Markdown version
    out_md = r"c:\Penilaian_Kinerja_Guru\Bab4_Perancangan_Antarmuka.md"
    with open(out_md, "w", encoding="utf-8") as f:
        f.write("# BAB IV: HASIL DAN PEMBAHASAN\n")
        f.write("## 4.4 Perancangan Antarmuka Pengguna (User Interface Design)\n\n")
        f.write("Perancangan antarmuka pengguna pada Sistem Pendukung Keputusan (SPK) Penilaian Kinerja Guru Metode MOORA SMA Al-Ihsan Boarding School disusun dalam bentuk rancangan wireframe monokromatik (polos perancangan) yang identik dengan sistem aktual:\n\n")
        
        for item in ui_data:
            f.write(f"### 4.4.{item['num']} {item['title']}\n\n")
            f.write(f"{item['fungsi']} **Hak Akses**: `{item['hak_akses']}`.\n\n")
            f.write(f"![{item['title']}](wireframes/{os.path.basename(item['img'])})\n\n")
            f.write(f"**Gambar 4.{item['num']}. {item['title']}**\n\n")
            f.write("| No | Nama Kontrol / Komponen | Tipe Kontrol | Fungsi / Keterangan |\n")
            f.write("| :---: | :--- | :---: | :--- |\n")
            for c in item['controls']:
                f.write(f"| {c[0]} | **{c[1]}** | `{c[2]}` | {c[3]} |\n")
            f.write("\n---\n\n")
            
    print(f"[SUCCESS] File Markdown berhasil disimpan di: {out_md}")

if __name__ == "__main__":
    generate_antarmuka_document()
