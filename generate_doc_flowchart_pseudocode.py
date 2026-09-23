import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="B0B0B0", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_image_caption(paragraph, num_suffix, title_text):
    """
    Menambahkan caption Word otomatis dengan SEQ field Gambar untuk Daftar Gambar
    """
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.keep_with_next = True
    
    r1 = paragraph.add_run("Gambar 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    fld_xml = r'''
        <w:fldSimple %s w:instr="SEQ Gambar \* ARABIC">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:t>%s</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), num_suffix)
    paragraph._p.append(parse_xml(fld_xml))
    
    r2 = paragraph.add_run(f" {title_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.bold = True

def add_table_caption(paragraph, num_suffix, title_text):
    """
    Menambahkan caption Word otomatis dengan SEQ field Tabel untuk Daftar Tabel
    """
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    paragraph.paragraph_format.space_before = Pt(10)
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.paragraph_format.keep_with_next = True
    
    r1 = paragraph.add_run("Tabel 4.")
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(11)
    r1.font.bold = True
    
    fld_xml = r'''
        <w:fldSimple %s w:instr="SEQ Tabel \* ARABIC">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="22"/>
                </w:rPr>
                <w:t>%s</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), num_suffix)
    paragraph._p.append(parse_xml(fld_xml))
    
    r2 = paragraph.add_run(f" {title_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(11)
    r2.font.bold = True

def generate_flowchart_pseudocode_doc():
    doc = docx.Document()
    
    # Page setup - A4, Margin Skripsi (Top 3cm, Left 4cm, Bottom 3cm, Right 3cm)
    for section in doc.sections:
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(3.0)
        section.bottom_margin = Cm(3.0)
        section.left_margin = Cm(4.0)
        section.right_margin = Cm(3.0)

    # Styles
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(12)
    style_normal.paragraph_format.line_spacing = 1.5
    style_normal.paragraph_format.space_after = Pt(6)

    # Header Bab IV
    h1 = doc.add_paragraph()
    h1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_h1 = h1.add_run("BAB IV\nHASIL DAN PEMBAHASAN\n4.3 PSEUDOCODE / FLOWCHART ALGORITMA MOORA")
    r_h1.font.name = 'Times New Roman'
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    h1.paragraph_format.space_after = Pt(18)

    # Pengantar
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_intro.add_run(
        "Penerapan metode Multi-Objective Optimization on the basis of Ratio Analysis (MOORA) pada Sistem Pendukung Keputusan "
        "Penilaian Kinerja Guru di SMA Al-Ihsan Boarding School dirancang secara sistematis ke dalam diagram alir (Flowchart) "
        "dan notasi algoritma terstruktur (Pseudocode). Perancangan ini menjadi acuan utama bagi pengembang sistem dalam "
        "mengimplementasikan logika pemrosesan data penilaian, perhitungan matriks keputusan, normalisasi rasio vektor, "
        "pembobotan preferensi, hingga penentuan ranking kinerja guru secara transparan dan terotomatisasi."
    )

    # 4.3.1 Flowchart Algoritma MOORA
    p_fc = doc.add_paragraph()
    p_fc.paragraph_format.line_spacing = 1.5
    p_fc.paragraph_format.space_before = Pt(12)
    p_fc.paragraph_format.space_after = Pt(4)
    p_fc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_fc_title = p_fc.add_run("4.3.1 Flowchart Algoritma MOORA\n")
    r_fc_title.font.bold = True
    r_fc_title.font.size = Pt(12)
    p_fc.add_run(
        "Flowchart algoritma MOORA menggambarkan alur logika komputasi yang berlangsung dari tahap inisialisasi data master "
        "hingga tahap penyimpanan hasil perankingan akhir ke dalam basis data. Diagram alir proses perhitungan MOORA "
        "disajikan pada gambar di bawah ini:"
    )

    # Embed Flowchart Image
    fc_img_path = "wireframes/flowchart_moora.png"
    if os.path.exists(fc_img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(8)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.paragraph_format.keep_with_next = True
        
        run_img = p_img.add_run()
        run_img.add_picture(fc_img_path, width=Cm(14.5))
        
        # Caption Word SEQ Gambar
        p_cap_fc = doc.add_paragraph(style='Caption')
        add_image_caption(p_cap_fc, "1", "Flowchart Proses Komputasi Algoritma Metode MOORA")

    # Penjelasan Alur Flowchart
    p_fc_desc = doc.add_paragraph()
    p_fc_desc.paragraph_format.line_spacing = 1.5
    p_fc_desc.paragraph_format.space_before = Pt(6)
    p_fc_desc.paragraph_format.space_after = Pt(4)
    p_fc_desc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_fc_desc.add_run(
        "Berdasarkan Gambar 4.1 di atas, alur eksekusi algoritma MOORA dapat diuraikan dalam langkah-langkah kerja berikut:\n"
        "1. Mulai (Start): Sistem menginisialisasi sesi eksekusi perhitungan untuk periode penilaian aktif yang dipilih oleh pengguna.\n"
        "2. Input Data: Sistem mengambil data master guru aktif (m alternatif), data kriteria penilaian aktif (n kriteria) beserta bobot preferensi (W), dan seluruh rekaman skor evaluasi kinerja guru pada periode tersebut.\n"
        "3. Validasi Kelengkapan Nilai: Sistem memeriksa apakah seluruh guru aktif telah dinilai secara lengkap pada seluruh kriteria. Jika terdapat guru yang belum dinilai, sistem menampilkan pesan peringatan dan membatalkan kalkulasi. Jika lengkap, alur berlanjut ke tahap kalkulasi.\n"
        "4. Pembentukan Matriks Keputusan (X): Menyusun matriks evaluasi kinerja guru berukuran m x n, di mana setiap elemen x_ij mewakili nilai guru ke-i pada kriteria ke-j.\n"
        "5. Perhitungan Pembagi Normalisasi: Menghitung nilai penyebut Euclidean untuk setiap kolom kriteria ke-j melalui akar dari akumulasi kuadrat nilai seluruh alternatif: |X_j| = akar(sum(x_ij^2)).\n"
        "6. Normalisasi Matriks (X*): Menghitung matriks ternormalisasi x*_ij = x_ij / |X_j| sehingga seluruh nilai terkonversi ke dalam rentang tanpa satuan.\n"
        "7. Pembobotan Matriks (V): Mengalikan setiap elemen matriks ternormalisasi dengan bobot kriteria yang bersesuaian: v_ij = w_j * x*_ij.\n"
        "8. Perhitungan Nilai Preferensi (Yi): Menghitung nilai optimasi multiobjektif setiap alternatif dengan menjumlahkan kriteria bertipe Benefit dan mengurangkannya dengan kriteria bertipe Cost: Yi = sum(Benefit) - sum(Cost).\n"
        "9. Perankingan Alternatif: Mengurutkan seluruh alternatif guru secara menurun (descending) berdasarkan nilai preferensi Yi terbesar sebagai peringkat terbaik.\n"
        "10. Penyimpanan & Output: Menyimpan hasil ranking dan detail matriks ke tabel basis data (HasilMoora & DetailMoora) serta menampilkan hasil perankingan ke antarmuka pengguna sistem.\n"
        "11. Selesai (End): Proses perhitungan SPK metode MOORA berhasil diselesaikan."
    )

    # 4.3.2 Pseudocode Algoritma MOORA
    p_pc = doc.add_paragraph()
    p_pc.paragraph_format.line_spacing = 1.5
    p_pc.paragraph_format.space_before = Pt(14)
    p_pc.paragraph_format.space_after = Pt(4)
    p_pc.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_pc_title = p_pc.add_run("4.3.2 Pseudocode Algoritma MOORA\n")
    r_pc_title.font.bold = True
    r_pc_title.font.size = Pt(12)
    p_pc.add_run(
        "Pseudocode algoritma MOORA menyajikan representasi tekstual dari logika prosedural yang diimplementasikan pada fungsi layanan "
        "kalkulasi sistem (layanan-moora.ts). Notasi algoritma dirancang menggunakan struktur pemrograman modular sebagai berikut:"
    )

    # Pseudocode Table Box (Bahasa Indonesia Akademik & Mudah Dipahami)
    pseudocode_lines = [
        "ALGORITMA Perhitungan_Metode_MOORA",
        "",
        "MASUKAN (INPUT):",
        "  - Data Guru Alternatif (A1, A2, ..., Am)",
        "  - Data Kriteria Penilaian (C1, C2, ..., Cn)",
        "  - Bobot Preferensi Kriteria (w1, w2, ..., wn)",
        "  - Sifat Atribut Kriteria (Benefit / Cost)",
        "  - Nilai Skor Evaluasi Kinerja Guru pada Periode Aktif",
        "",
        "KELUARAN (OUTPUT):",
        "  - Nilai Preferensi Multiobjektif (Yi) Setiap Guru",
        "  - Daftar Urutan Peringkat / Ranking Kinerja Guru",
        "",
        "LANGKAH-LANGKAH PROSEDURAL:",
        "1.   MULAI",
        "2.   Baca Data Guru (m), Kriteria (n), Bobot Kriteria (W), dan Skor Penilaian",
        "3.   ",
        "4.   // Tahap 1: Validasi Kelengkapan Nilai",
        "5.   Periksa apakah semua guru telah memiliki nilai lengkap pada seluruh kriteria",
        "6.   JIKA data nilai belum lengkap MAKA",
        "7.       Tampilkan pesan 'Data penilaian guru belum lengkap'",
        "8.       Hentikan proses dan KELUAR",
        "9.   AKHIR-JIKA",
        "10.  ",
        "11.  // Tahap 2: Pembentukan Matriks Keputusan (X)",
        "12.  UNTUK setiap guru (i = 1 sampai m) LAKUKAN",
        "13.      UNTUK setiap kriteria (j = 1 sampai n) LAKUKAN",
        "14.          X[i][j] <- Nilai skor evaluasi guru ke-i pada kriteria ke-j",
        "15.      AKHIR-UNTUK",
        "16.  AKHIR-UNTUK",
        "17.  ",
        "18.  // Tahap 3: Hitung Pembagi Normalisasi Vektor Euclidean",
        "19.  UNTUK setiap kriteria (j = 1 sampai n) LAKUKAN",
        "20.      TotalKuadrat <- 0",
        "21.      UNTUK setiap guru (i = 1 sampai m) LAKUKAN",
        "22.          TotalKuadrat <- TotalKuadrat + (X[i][j])^2",
        "23.      AKHIR-UNTUK",
        "24.      Pembagi[j] <- AKAR_KUADRAT(TotalKuadrat)",
        "25.  AKHIR-UNTUK",
        "26.  ",
        "27.  // Tahap 4: Normalisasi dan Pembobotan Matriks",
        "28.  UNTUK setiap guru (i = 1 sampai m) LAKUKAN",
        "29.      UNTUK setiap kriteria (j = 1 sampai n) LAKUKAN",
        "30.          // Normalisasi nilai: x*_ij = x_ij / Pembagi_j",
        "31.          X_Normalisasi[i][j] <- X[i][j] / Pembagi[j]",
        "32.          // Pembobotan nilai: v_ij = Bobot_j * x*_ij",
        "33.          V[i][j] <- Bobot[j] * X_Normalisasi[i][j]",
        "34.      AKHIR-UNTUK",
        "35.  AKHIR-UNTUK",
        "36.  ",
        "37.  // Tahap 5: Perhitungan Nilai Preferensi Multiobjektif (Yi)",
        "38.  UNTUK setiap guru (i = 1 sampai m) LAKUKAN",
        "39.      TotalBenefit <- 0",
        "40.      TotalCost    <- 0",
        "41.      UNTUK setiap kriteria (j = 1 sampai n) LAKUKAN",
        "42.          JIKA Kriteria[j] adalah BENEFIT MAKA",
        "43.              TotalBenefit <- TotalBenefit + V[i][j]",
        "44.          JIKA_TIDAK",
        "45.              TotalCost    <- TotalCost + V[i][j]",
        "46.          AKHIR-JIKA",
        "47.      AKHIR-UNTUK",
        "48.      // Nilai Yi = Total Benefit - Total Cost",
        "49.      NilaiPreferensi[i] <- TotalBenefit - TotalCost",
        "50.  AKHIR-UNTUK",
        "51.  ",
        "52.  // Tahap 6: Penentuan Ranking / Urutan Peringkat",
        "53.  Urutkan daftar guru berdasarkan NilaiPreferensi secara MENURUN (terbesar ke terkecil)",
        "54.  UNTUK setiap peringkat (k = 1 sampai m) LAKUKAN",
        "55.      GuruTerurut[k].Ranking <- k",
        "56.  AKHIR-UNTUK",
        "57.  ",
        "58.  // Tahap 7: Penyimpanan dan Penampilan Hasil",
        "59.  Simpan hasil perhitungan Nilai Preferensi dan Ranking ke dalam basis data",
        "60.  Tampilkan tabel hasil perankingan kinerja guru",
        "61.  SELESAI"
    ]

    tbl_pc = doc.add_table(rows=1, cols=1)
    tbl_pc.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_pc.autofit = False
    tbl_pc.columns[0].width = Cm(15.0)
    set_table_borders(tbl_pc, color="4B5563", sz="6", val="single")

    cell_pc = tbl_pc.cell(0, 0)
    set_cell_background(cell_pc, "F8FAFC")
    set_cell_margins(cell_pc, top=140, bottom=140, left=180, right=180)

    # Insert pseudocode lines
    p_cell = cell_pc.paragraphs[0]
    p_cell.paragraph_format.line_spacing = 1.15
    p_cell.paragraph_format.space_before = Pt(0)
    p_cell.paragraph_format.space_after = Pt(0)

    for idx, line in enumerate(pseudocode_lines):
        if idx > 0:
            p_cell = cell_pc.add_paragraph()
            p_cell.paragraph_format.line_spacing = 1.15
            p_cell.paragraph_format.space_before = Pt(0)
            p_cell.paragraph_format.space_after = Pt(0)

        run = p_cell.add_run(line)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        
        # Color & highlight keywords
        if any(line.startswith(k) for k in ["ALGORITMA", "MASUKAN", "KELUARAN", "LANGKAH-LANGKAH", "SELESAI"]):
            run.font.bold = True
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif any(k in line for k in ["// Tahap"]):
            run.font.bold = True
            run.font.color.rgb = RGBColor(37, 99, 235)
        elif line.strip().startswith("//"):
            run.font.italic = True
            run.font.color.rgb = RGBColor(100, 116, 139)

    # Analisis Kompleksitas Algoritma
    p_comp = doc.add_paragraph()
    p_comp.paragraph_format.line_spacing = 1.5
    p_comp.paragraph_format.space_before = Pt(14)
    p_comp.paragraph_format.space_after = Pt(4)
    p_comp.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_comp_title = p_comp.add_run("4.3.3 Analisis Efisiensi dan Kompleksitas Algoritma MOORA\n")
    r_comp_title.font.bold = True
    r_comp_title.font.size = Pt(12)
    p_comp.add_run(
        "Berdasarkan pseudocode yang dirancang di atas, efisiensi kinerja algoritma MOORA dapat dianalisis dari dua aspek utama:\n"
        "1. Kompleksitas Waktu (Time Complexity): Tahap pembentukan matriks, normalisasi, dan pembobotan memerlukan waktu komputasi linier terhadap perkalian jumlah guru (m) dan kriteria (n), yaitu O(m x n). Tahap perankingan akhir menggunakan algoritma pengurutan cepat (QuickSort/TimSort) dengan kompleksitas O(m log m). Sehingga total kompleksitas waktu algoritma adalah O(m x n + m log m). Pada data 18 guru dan 4 kriteria, kalkulasi dapat diselesaikan dalam waktu kurang dari 50 milidetik (< 0.05 detik).\n"
        "2. Kompleksitas Ruang (Space Complexity): Ruang penyimpanan memori yang dibutuhkan untuk mengalokasikan matriks keputusan, matriks normalisasi, dan matriks terbobot adalah berorde O(m x n), yang sangat efisien dan ringan untuk dijalankan pada arsitektur server web modern."
    )

    # Save to DOCX
    out_docx = r"c:\Penilaian_Kinerja_Guru\Bab4_Flowchart_Pseudocode_MOORA.docx"
    try:
        doc.save(out_docx)
        print(f"[SUCCESS] File DOCX berhasil disimpan di: {out_docx}")
    except PermissionError:
        alt_docx = r"c:\Penilaian_Kinerja_Guru\Bab4_Flowchart_Pseudocode_MOORA_Baru.docx"
        doc.save(alt_docx)
        print(f"[WARNING] File Bab4_Flowchart_Pseudocode_MOORA.docx sedang dibuka di Word.")
        print(f"[SUCCESS] Disimpan sebagai: {alt_docx}")

    # Generate Markdown version
    out_md = r"c:\Penilaian_Kinerja_Guru\Bab4_Flowchart_Pseudocode_MOORA.md"
    with open(out_md, "w", encoding="utf-8") as f:
        f.write("# BAB IV: HASIL DAN PEMBAHASAN\n")
        f.write("## 4.3 Pseudocode / Flowchart Algoritma MOORA\n\n")
        f.write("Penerapan metode MOORA pada Sistem Pendukung Keputusan Penilaian Kinerja Guru di SMA Al-Ihsan Boarding School dirancang ke dalam diagram alir (Flowchart) dan notasi algoritma terstruktur (Pseudocode).\n\n")
        
        f.write("### 4.3.1 Flowchart Algoritma MOORA\n\n")
        f.write("![Flowchart Algoritma MOORA](wireframes/flowchart_moora.png)\n\n")
        f.write("**Gambar 4.1. Flowchart Proses Komputasi Algoritma Metode MOORA**\n\n")
        f.write("#### Uraian Langkah Alur Flowchart:\n")
        f.write("1. **Mulai (Start)**: Sistem menginisialisasi sesi eksekusi perhitungan periode aktif.\n")
        f.write("2. **Input Data**: Mengambil data master guru ($m$), kriteria ($n$), bobot preferensi ($W$), dan skor penilaian.\n")
        f.write("3. **Validasi Kelengkapan Nilai**: Memeriksa apakah seluruh guru telah dinilai lengkap. Jika tidak, proses dibatalkan.\n")
        f.write("4. **Pembentukan Matriks Keputusan ($X$)**: Menyusun matriks nilai $X = [x_{ij}]_{m \\times n}$.\n")
        f.write("5. **Perhitungan Pembagi Normalisasi**: Menghitung penyebut Euclidean $|X_j| = \\sqrt{\\sum_{i=1}^m (x_{ij})^2}$.\n")
        f.write("6. **Normalisasi Matriks ($X^*$)**: Menghitung rasio vektor $x^*_{ij} = \\frac{x_{ij}}{|X_j|}$.\n")
        f.write("7. **Pembobotan Matriks ($V$)**: Menghitung matriks terbobot $v_{ij} = w_j \\times x^*_{ij}$.\n")
        f.write("8. **Perhitungan Nilai Preferensi ($Y_i$)**: Menghitung $Y_i = \\sum \\text{Benefit} - \\sum \\text{Cost}$.\n")
        f.write("9. **Perankingan Alternatif**: Mengurutkan nilai $Y_i$ secara *descending* (terbesar ke terkecil).\n")
        f.write("10. **Penyimpanan & Output**: Menyimpan hasil ke database (`HasilMoora`) dan menampilkan leaderboard.\n")
        f.write("11. **Selesai (End)**: Eksekusi SPK MOORA selesai.\n\n")
        f.write("---\n\n")

        f.write("### 4.3.2 Pseudocode Algoritma MOORA\n\n")
        f.write("```pascal\n")
        for line in pseudocode_lines:
            f.write(f"{line}\n")
        f.write("```\n\n")
        f.write("---\n\n")

        f.write("### 4.3.3 Analisis Efisiensi dan Kompleksitas Algoritma MOORA\n\n")
        f.write("- **Kompleksitas Waktu (Time Complexity)**: $\\mathcal{O}(m \\times n + m \\log m)$, di mana $m$ adalah jumlah alternatif guru dan $n$ adalah jumlah kriteria. Waktu eksekusi sangat cepat ($< 0.05$ detik).\n")
        f.write("- **Kompleksitas Ruang (Space Complexity)**: $\\mathcal{O}(m \\times n)$, sangat hemat alokasi memori server.\n")

    print(f"[SUCCESS] File Markdown berhasil disimpan di: {out_md}")

if __name__ == "__main__":
    generate_flowchart_pseudocode_doc()
