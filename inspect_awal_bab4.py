import docx
from lxml import etree
from docx.oxml.ns import qn

doc = docx.Document('SKRIPSI.docx')
body_elements = list(doc._body._body)

print(f"Total body elements: {len(body_elements)}")
print()

# Cari semua paragraf yang mengandung "Tabel 4.1", "Tabel 4.2", "Tabel 4.3", "4.2.2", "4.1.1"
print("=== Scanning for Tabel 4.1, 4.2, 4.3 and surrounding elements ===")
found_indices = []

for i, elem in enumerate(body_elements):
    tag = elem.tag.split('}')[-1]
    if tag == 'p':
        xml_bytes = etree.tostring(elem)
        xml_str = xml_bytes.decode('utf-8', errors='ignore')
        text = ''.join(t.text or '' for t in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
        
        # Check for relevant content
        check_terms = ['Tabel 4.1', 'Tabel 4.2', 'Tabel 4.3', '4.2.2', '4.1.1', 
                       'Kebutuhan Fungsional', 'Kebutuhan Non', 'Aktor', 'BAB IV',
                       'SEQ Tabel']
        for term in check_terms:
            if term in xml_str or term in text:
                found_indices.append(i)
                break

print(f"Found {len(found_indices)} relevant elements")
print()

# Show context around each found element
shown = set()
for idx in found_indices:
    start = max(0, idx - 2)
    end = min(len(body_elements), idx + 4)
    
    if idx in shown:
        continue
    
    for j in range(start, end):
        if j in shown:
            continue
        shown.add(j)
        
        elem = body_elements[j]
        tag = elem.tag.split('}')[-1]
        
        if tag == 'p':
            text = ''.join(t.text or '' for t in elem.iter('{http://schemas.openxmlformats.org/wordprocessingml/2006/main}t'))
            xml_bytes = etree.tostring(elem)
            xml_str = xml_bytes.decode('utf-8', errors='ignore')
            
            # Get style
            pPr = elem.find(qn('w:pPr'))
            style = ''
            if pPr is not None:
                pStyle = pPr.find(qn('w:pStyle'))
                if pStyle is not None:
                    style = pStyle.get(qn('w:val'), '')
            
            has_seq = 'SEQ Tabel' in xml_str
            has_section = 'sectPr' in xml_str
            
            print(f"  [{j}] P | style={style} | SEQ={has_seq} | sectPr={has_section}")
            print(f"       text={repr(text[:120])}")
        elif tag == 'tbl':
            rows = elem.findall('.//' + qn('w:tr'))
            print(f"  [{j}] TBL | rows={len(rows)}")
        elif tag == 'sectPr':
            print(f"  [{j}] SECTPR")
        else:
            print(f"  [{j}] {tag.upper()}")
    
    print("  ---")

print()
print("=== Check paragraphs doc.paragraphs around Tabel 4.1/4.2/4.3 ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if any(x in t for x in ['Tabel 4.1', 'Tabel 4.2', 'Tabel 4.3', '4.1.1', '4.2.2', 
                              'Kebutuhan Fungsional', 'Kebutuhan Non', 'Aktor', 'analisis']):
        print(f"  Para[{i}] style={p.style.name!r} | {repr(t[:100])}")
