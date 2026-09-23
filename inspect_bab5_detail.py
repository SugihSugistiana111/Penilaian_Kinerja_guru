import docx

doc = docx.Document('BAB V.docx')
print(f"Total paragraphs: {len(doc.paragraphs)}")
for i, p in enumerate(doc.paragraphs):
    runs_info = [f"'{r.text}' (b={r.bold}, i={r.italic})" for r in p.runs]
    print(f"P{i} [{p.style.name}]: text='{p.text}' | runs={runs_info}")

print(f"\nTotal tables: {len(doc.tables)}")
for i, tbl in enumerate(doc.tables):
    print(f"Table {i}: {len(tbl.rows)} rows, {len(tbl.columns)} cols")
    for r_idx, r in enumerate(tbl.rows):
        print(f"  R{r_idx}: {[c.text.replace(chr(10), ' ') for c in r.cells]}")
