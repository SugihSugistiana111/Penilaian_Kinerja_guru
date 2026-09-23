import docx

doc = docx.Document('SKRIPSI.docx')

# Find BAB IV content range
bab4_start = None
bab4_end = len(doc.paragraphs)

for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if 'BAB IV' in t and bab4_start is None:
        bab4_start = i
    if ('BAB V' in t or 'DAFTAR PUSTAKA' in t) and bab4_start is not None and i > bab4_start + 5:
        bab4_end = i
        break

print(f"BAB IV range: P{bab4_start} to P{bab4_end}")
print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")
print()

# Show ALL paragraphs containing "Tabel" or "Gambar" in the entire document
print("=== ALL Tabel/Gambar caption paragraphs ===")
for i, p in enumerate(doc.paragraphs):
    t = p.text.strip()
    if t.startswith('Tabel ') or t.startswith('Gambar '):
        has_img = 'w:drawing' in p._p.xml or 'w:pict' in p._p.xml
        has_fld = 'w:fldSimple' in p._p.xml or 'w:instrText' in p._p.xml
        print(f"P{i}: [{p.style.name}] \"{t[:100]}\" | HasImage:{has_img} | HasSEQ:{has_fld}")

print()
print("=== Paragraphs with images (w:drawing) around BAB IV area ===")
for i in range(700, min(len(doc.paragraphs), 1200)):
    p = doc.paragraphs[i]
    has_img = 'w:drawing' in p._p.xml
    t = p.text.strip()
    if has_img or t.startswith('Tabel') or t.startswith('Gambar') or t.startswith('AD-') or t.startswith('SD-'):
        print(f"P{i}: \"{t[:80]}\" | HasImage:{has_img}")
