import docx
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = docx.Document('SKRIPSI.docx')

# Find Tabel 4.5 caption (SEQ Tabel #5 in BAB IV context)
# From earlier inspection, the Use Case tables start around P764+
# Let's find all "Tabel" SEQ captions and their surrounding context

tabel_seq_count = 0
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    if 'SEQ Tabel' in xml:
        tabel_seq_count += 1
        if tabel_seq_count >= 3 and tabel_seq_count <= 8:
            t = p.text.strip()
            sb = p.paragraph_format.space_before
            sa = p.paragraph_format.space_after
            print(f"Tabel SEQ #{tabel_seq_count} at P{i}: '{t[:80]}'")
            print(f"  space_before={sb}, space_after={sa}")
            
            # Check paragraphs around it
            for offset in range(-3, 4):
                j = i + offset
                if 0 <= j < len(doc.paragraphs):
                    pp = doc.paragraphs[j]
                    has_img = 'w:drawing' in pp._p.xml
                    sb2 = pp.paragraph_format.space_before
                    sa2 = pp.paragraph_format.space_after
                    print(f"  P{j} [{pp.style.name}]: '{pp.text[:60]}' | sb={sb2} sa={sa2} img={has_img}")
            print()

# Now check tables near Tabel 4.5
print("=== Tables in document ===")
for idx, tbl in enumerate(doc.tables):
    # Check first cell text
    first_cell = tbl.cell(0, 0).text[:40] if tbl.rows else "?"
    num_rows = len(tbl.rows)
    
    # Check if repeat header is set
    first_row = tbl.rows[0]
    trPr = first_row._tr.get_or_add_trPr()
    tblHeader = trPr.find(qn('w:tblHeader'))
    has_repeat = tblHeader is not None
    
    if idx >= 2 and idx <= 10:
        print(f"Table #{idx}: rows={num_rows}, firstCell='{first_cell}', repeatHeader={has_repeat}")
