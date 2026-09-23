import docx
from lxml import etree

doc = docx.Document('SKRIPSI_Updated.docx')

print("Inspect SKRIPSI_Updated.docx around AD-11 and SD-11:")
for i, p in enumerate(doc.paragraphs):
    txt = p.text.strip()
    xml_s = p._p.xml
    has_img = 'w:drawing' in xml_s or 'w:pict' in xml_s
    if any(k in txt for k in ['AD-10', 'AD-11', 'AD-12', 'SD-10', 'SD-11', 'SD-12']):
        print(f"P{i} (img={has_img}): '{txt}'")
        for off in range(1, 3):
            if i + off < len(doc.paragraphs):
                p_next = doc.paragraphs[i+off]
                has_img_next = 'w:drawing' in p_next._p.xml or 'w:pict' in p_next._p.xml
                print(f"   + {off} P{i+off} (img={has_img_next}): '{p_next.text.strip()[:60]}'")
