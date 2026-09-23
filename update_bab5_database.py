import os
import docx
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_background(cell, color_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:insideV w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:left w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
            <w:right w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def add_caption_field(p, label_prefix, seq_num, desc_text, seq_type="Gambar"):
    """
    Menambahkan caption Word otomatis dengan SEQ field yang kompatibel 
    dengan References -> Insert Table of Figures (Daftar Gambar / Daftar Tabel).
    """
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.keep_with_next = False
    
    # Run 1: Label prefix (Bold, Times New Roman 12pt)
    r1 = p.add_run(label_prefix)
    r1.font.name = 'Times New Roman'
    r1.font.size = Pt(12)
    r1.font.bold = True
    
    # SEQ Field
    fld_xml = '''
        <w:fldSimple %s w:instr=" SEQ %s \\* ARABIC ">
            <w:r>
                <w:rPr>
                    <w:rFonts w:ascii="Times New Roman" w:hAnsi="Times New Roman"/>
                    <w:b/>
                    <w:sz w:val="24"/>
                </w:rPr>
                <w:t>%d</w:t>
            </w:r>
        </w:fldSimple>
    ''' % (nsdecls("w"), seq_type, seq_num)
    p._p.append(parse_xml(fld_xml))
    
    # Run 2: Description text (Regular, Times New Roman 12pt)
    r2 = p.add_run(f" {desc_text}")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)
    r2.font.bold = False

def create_full_bab5_doc():
    doc = docx.Document('BAB V.docx')
    
    # Set page margins: Top 3cm, Bottom 3cm, Left 4cm, Right 3cm
    sections = doc.sections
    for s in sections:
        s.top_margin = Cm(3)
        s.bottom_margin = Cm(3)
        s.left_margin = Cm(4)
        s.right_margin = Cm(3)
        
    new_doc = docx.Document('BAB V.docx')
    
    # Clear all existing body elements in new_doc
    body = new_doc._body._body
    for child in list(body):
        tag = child.tag.split('}')[-1]
        if tag not in ('sectPr',):
            body.remove(child)
            
    # Function to add heading 1
    def add_h1(text1, text2):
        p1 = new_doc.add_paragraph(text1, style='Heading 1')
        p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(0)
        p1.paragraph_format.line_spacing = 1.5
        for r in p1.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = True
            
        p2 = new_doc.add_paragraph(text2, style='Heading 1')
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.paragraph_format.space_before = Pt(0)
        p2.paragraph_format.space_after = Pt(18)
        p2.paragraph_format.line_spacing = 1.5
        for r in p2.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(14)
            r.font.bold = True

    def add_h2(text):
        p = new_doc.add_paragraph(style='Heading 2')
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.keep_with_next = True
        
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True
        
        # Add numPr to Heading 2 to match original style
        pPr = p._p.get_or_add_pPr()
        numPr = parse_xml(f'<w:numPr {nsdecls("w")}><w:ilvl w:val="1"/><w:numId w:val="1"/></w:numPr>')
        pPr.append(numPr)
        return p

    def add_h3(num_str, title):
        p = new_doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        p.paragraph_format.keep_with_next = True
        
        r = p.add_run(f"{num_str}\t{title}")
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.font.bold = True
        return p

    def add_p(text, indent=0.75):
        p = new_doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        if indent > 0:
            p.paragraph_format.first_line_indent = Cm(indent)
            
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        return p

    def add_image(img_path, width_cm=15.2):
        p = new_doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        
        r = p.add_run()
        r.add_picture(img_path, width=Cm(width_cm))
        return p

    # --- Start building Chapter V ---
    add_h1("BAB V", "IMPLEMENTASI DAN PENGUJIAN SISTEM")
    
    # 5.1 Implementasi Perangkat Keras
    add_h2("Implementasi Perangkat Keras")
    add_p("Implementasi perangkat keras merupakan tahap yang menjelaskan perangkat yang digunakan untuk menjalankan dan menguji sistem pendukung keputusan penilaian kinerja guru. Perangkat yang digunakan berupa komputer laptop dengan spesifikasi yang mendukung proses pengembangan dan pengujian sistem.")
    
    # Table 5.1 Caption
    p_cap_t1 = new_doc.add_paragraph()
    add_caption_field(p_cap_t1, "Tabel 5. ", 1, "Spesifikasi Perangkat Keras", seq_type="Tabel")
    p_cap_t1.paragraph_format.space_before = Pt(6)
    p_cap_t1.paragraph_format.space_after = Pt(4)
    p_cap_t1.paragraph_format.keep_with_next = True
    
    # Table 5.1
    t1_data = [
        ['No', 'Komponen', 'Spesifikasi'],
        ['1', 'Perangkat', 'HP Notebook 14-an002ax / Laptop Development'],
        ['2', 'Prosesor', 'AMD A8-7410 APU with AMD Radeon R5 Graphics, 2,20 GHz'],
        ['3', 'Memori (RAM)', '8 GB DDR3L'],
        ['4', 'Grafis', 'AMD Radeon R5 Graphics'],
        ['5', 'Penyimpanan', 'SSD 256 GB SATA III']
    ]
    tbl1 = new_doc.add_table(rows=len(t1_data), cols=3)
    tbl1.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl1, color="CCCCCC", sz="4", val="single")
    col_w1 = [Cm(1.2), Cm(4.0), Cm(9.8)]
    for r_idx, row in enumerate(tbl1.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_w1[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            cell_p = cell.paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            cell_p.paragraph_format.space_before = Pt(2)
            cell_p.paragraph_format.space_after = Pt(2)
            cell_p.paragraph_format.line_spacing = 1.15
            run = cell_p.add_run(t1_data[r_idx][c_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if r_idx == 0:
                set_cell_background(cell, "F2F2F2")
                run.font.bold = True
                
    add_p("Perangkat keras tersebut digunakan sebagai perangkat utama dalam proses pengembangan, pengujian, dan menjalankan sistem. Spesifikasi prosesor dan RAM yang digunakan telah mendukung kebutuhan sistem selama proses implementasi dan pengujian.")
    
    # 5.2 Implementasi Perangkat Lunak
    add_h2("Implementasi Perangkat Lunak")
    add_p("Implementasi perangkat lunak menjelaskan perangkat lunak yang digunakan dalam proses pengembangan, pengujian, dan menjalankan Sistem Pendukung Keputusan Penilaian Kinerja Guru. Perangkat lunak yang digunakan disesuaikan dengan kebutuhan pengembangan sistem.")
    
    # Table 5.2 Caption
    p_cap_t2 = new_doc.add_paragraph()
    add_caption_field(p_cap_t2, "Tabel 5. ", 2, "Spesifikasi Perangkat Lunak", seq_type="Tabel")
    p_cap_t2.paragraph_format.space_before = Pt(6)
    p_cap_t2.paragraph_format.space_after = Pt(4)
    p_cap_t2.paragraph_format.keep_with_next = True
    
    # Table 5.2
    t2_data = [
        ['No', 'Perangkat Lunak', 'Keterangan'],
        ['1', 'Windows 11 Home 64-bit', 'Sistem operasi yang digunakan pada perangkat pengembangan.'],
        ['2', 'Visual Studio Code', 'Text editor yang digunakan untuk menulis dan mengelola kode program.'],
        ['3', 'Node.js (v20.x)', 'JavaScript runtime environment untuk menjalankan aplikasi Next.js.'],
        ['4', 'Next.js 14 (App Router)', 'Framework web full-stack React untuk membangun antarmuka dan backend API.'],
        ['5', 'TypeScript', 'Bahasa pemrograman strongly typed untuk pengembangan logika sistem.'],
        ['6', 'Prisma ORM (v5.21)', 'Object-Relational Mapping untuk interaksi query dan migrasi basis data.'],
        ['7', 'SQLite 3', 'Relational Database Management System (RDBMS) penyimpanan data sistem.'],
        ['8', 'Tailwind CSS', 'Utility-first CSS framework untuk perancangan antarmuka responsif.'],
        ['9', 'Google Chrome Browser', 'Peramban web untuk mengakses, menguji, dan memverifikasi antarmuka sistem.'],
        ['10', 'Git & GitHub', 'Version control system untuk manajemen dan riwayat versi kode program.']
    ]
    tbl2 = new_doc.add_table(rows=len(t2_data), cols=3)
    tbl2.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl2, color="CCCCCC", sz="4", val="single")
    col_w2 = [Cm(1.2), Cm(4.8), Cm(9.0)]
    for r_idx, row in enumerate(tbl2.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if r_idx == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for c_idx, cell in enumerate(row.cells):
            cell.width = col_w2[c_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
            cell_p = cell.paragraphs[0]
            cell_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx == 0 or r_idx == 0 else WD_ALIGN_PARAGRAPH.LEFT
            cell_p.paragraph_format.space_before = Pt(2)
            cell_p.paragraph_format.space_after = Pt(2)
            cell_p.paragraph_format.line_spacing = 1.15
            run = cell_p.add_run(t2_data[r_idx][c_idx])
            run.font.name = 'Times New Roman'
            run.font.size = Pt(11)
            if r_idx == 0:
                set_cell_background(cell, "F2F2F2")
                run.font.bold = True

    add_p("Perangkat lunak tersebut digunakan untuk mendukung proses pengembangan dan pengujian sistem. Setiap perangkat lunak memiliki fungsi yang berbeda, mulai dari penulisan kode, pengelolaan basis data, hingga menjalankan dan menguji sistem melalui peramban.")

    # 5.3 Implementasi Database
    add_h2("Implementasi Database")
    add_p("Implementasi basis data merupakan tahapan realisasi dari perancangan struktur tabel yang telah disusun pada Bab IV ke dalam sistem manajemen basis data SQLite yang dikelola menggunakan Prisma ORM. Basis data pada Sistem Pendukung Keputusan Penilaian Kinerja Guru terdiri dari 11 (sebelas) tabel utama yang saling berelasi untuk mengelola data tingkat hak akses pengguna, master profil guru, kriteria dan bobot penilaian, rubrik skala penilaian, periode penilaian berkala, transaksi penilaian kinerja, hasil kalkulasi metode MOORA, hingga pencatatan riwayat log aktivitas sistem. Berikut merupakan penjelasan singkat dan implementasi dari masing-masing tabel pada basis data:")

    # 11 Database Tables Configuration
    db_sections = [
        {
            "num": "5.3.1",
            "title": "Implementasi Tabel Role (tbl_role)",
            "desc": "Tabel Role (tbl_role) berfungsi untuk menyimpan data tingkat hak akses atau peran pengguna dalam sistem yang mencakup Administrator, Kepala Sekolah, dan Guru. Tabel ini memiliki atribut id bertipe CUID sebagai primary key dan atribut namaRole yang bersifat unik. Implementasi data pada tabel Role disajikan pada Gambar 5.1.",
            "img": "db_screenshots/img_db_01_role.png",
            "seq_num": 1,
            "cap_text": "Implementasi Tabel Role (tbl_role)"
        },
        {
            "num": "5.3.2",
            "title": "Implementasi Tabel User (tbl_user)",
            "desc": "Tabel User (tbl_user) digunakan untuk menyimpan data kredensial akun pengguna sistem guna keperluan autentikasi login dan otorisasi. Tabel ini menyimpan atribut id (primary key), nama, username (unik), password yang telah diamankan menggunakan algoritma hash Bcrypt, foreign key roleId yang berelasi ke tabel Role, foreign key guruId yang berelasi ke tabel Guru, serta status keaktifan akun. Implementasi data pada tabel User disajikan pada Gambar 5.2.",
            "img": "db_screenshots/img_db_02_user.png",
            "seq_num": 2,
            "cap_text": "Implementasi Tabel User (tbl_user)"
        },
        {
            "num": "5.3.3",
            "title": "Implementasi Tabel Guru (tbl_guru)",
            "desc": "Tabel Guru (tbl_guru) digunakan untuk mengelola data master profil pendidik yang menjadi alternatif dalam evaluasi penilaian kinerja di SMA Al-Ihsan Boarding School. Atribut yang dikelola meliputi id (primary key), NIP (unik), nama lengkap guru beserta gelar, jenis kelamin, jabatan tugas mengajar, tugas tambahan utama, tugas tambahan lain, jumlah siswa per rombel, jumlah jam ajar tatap muka, serta status keaktifan guru. Implementasi data pada tabel Guru disajikan pada Gambar 5.3.",
            "img": "db_screenshots/img_db_03_guru.png",
            "seq_num": 3,
            "cap_text": "Implementasi Tabel Guru (tbl_guru)"
        },
        {
            "num": "5.3.4",
            "title": "Implementasi Tabel Kriteria (tbl_kriteria)",
            "desc": "Tabel Kriteria (tbl_kriteria) berfungsi untuk menyimpan parameter kriteria evaluasi kinerja guru beserta konfigurasi bobot preferensi dan jenis atribut dalam metode MOORA. Kriteria yang dikelola meliputi Kehadiran (C1; bobot 0,35), Ketepatan Waktu (C2; bobot 0,25), Kelengkapan Perangkat Pembelajaran (C3; bobot 0,25), dan Kelengkapan Administrasi Penilaian (C4; bobot 0,15) dengan seluruh kriteria bertipe Benefit. Implementasi data pada tabel Kriteria disajikan pada Gambar 5.4.",
            "img": "db_screenshots/img_db_04_kriteria.png",
            "seq_num": 4,
            "cap_text": "Implementasi Tabel Kriteria (tbl_kriteria)"
        },
        {
            "num": "5.3.5",
            "title": "Implementasi Tabel Skala Penilaian (tbl_skala_penilaian)",
            "desc": "Tabel Skala Penilaian (tbl_skala_penilaian) digunakan untuk menyimpan rubrik tingkatan skor penilaian (sub-kriteria) yang menjadi panduan evaluasi kompetensi guru. Tabel ini menyimpan skala skor kuantitatif 1 sampai 5, label predikat (Sangat Kurang, Kurang, Cukup, Baik, hingga Sangat Baik), serta deskripsi indikator pencapaian kompetensi dengan foreign key kriteriaId yang berelasi ke tabel Kriteria. Implementasi data pada tabel Skala Penilaian disajikan pada Gambar 5.5.",
            "img": "db_screenshots/img_db_05_skala_penilaian.png",
            "seq_num": 5,
            "cap_text": "Implementasi Tabel Skala Penilaian (tbl_skala_penilaian)"
        },
        {
            "num": "5.3.6",
            "title": "Implementasi Tabel Periode Penilaian (tbl_periode_penilaian)",
            "desc": "Tabel Periode Penilaian (tbl_periode_penilaian) berfungsi untuk mengorganisir agenda pelaksanaan evaluasi kinerja guru secara berkala. Tabel ini memuat atribut bulan, tahun pelaksanaan, label nama periode (contoh: 'Mei 2026' dan 'Juni 2026'), serta status operasional evaluasi (DRAFT, AKTIF, atau SELESAI). Implementasi data pada tabel Periode Penilaian disajikan pada Gambar 5.6.",
            "img": "db_screenshots/img_db_06_periode_penilaian.png",
            "seq_num": 6,
            "cap_text": "Implementasi Tabel Periode Penilaian (tbl_periode_penilaian)"
        },
        {
            "num": "5.3.7",
            "title": "Implementasi Tabel Penilaian (tbl_penilaian)",
            "desc": "Tabel Penilaian (tbl_penilaian) digunakan untuk menyimpan transaksi evaluasi kinerja setiap guru pada periode tertentu. Tabel ini merekam relasi foreign key guruId ke tabel Guru, periodeId ke tabel Periode Penilaian, nama penilai/evaluator (Kepala Sekolah), status penilaian, serta catatan kualitatif hasil pembinaan. Implementasi data pada tabel Penilaian disajikan pada Gambar 5.7.",
            "img": "db_screenshots/img_db_07_penilaian.png",
            "seq_num": 7,
            "cap_text": "Implementasi Tabel Penilaian (tbl_penilaian)"
        },
        {
            "num": "5.3.8",
            "title": "Implementasi Tabel Detail Penilaian (tbl_detail_penilaian)",
            "desc": "Tabel Detail Penilaian (tbl_detail_penilaian) digunakan untuk menyimpan rincian skor nilai asli (raw score) yang diberikan penilai untuk masing-masing kriteria pada setiap entri penilaian guru (skala nilai 1,0 s.d. 5,0). Data pada tabel ini menjadi input pembentukan matriks keputusan awal (X) pada metode MOORA. Implementasi data pada tabel Detail Penilaian disajikan pada Gambar 5.8.",
            "img": "db_screenshots/img_db_08_detail_penilaian.png",
            "seq_num": 8,
            "cap_text": "Implementasi Tabel Detail Penilaian (tbl_detail_penilaian)"
        },
        {
            "num": "5.3.9",
            "title": "Implementasi Tabel Hasil MOORA (tbl_hasil_moora)",
            "desc": "Tabel Hasil MOORA (tbl_hasil_moora) digunakan untuk menyimpan nilai akhir preferensi multiobjektif (Yi) yang diperoleh dari hasil kalkulasi optimasi MOORA beserta urutan peringkat (ranking) kinerja guru pada periode penilaian terkait. Implementasi data pada tabel Hasil MOORA disajikan pada Gambar 5.9.",
            "img": "db_screenshots/img_db_09_hasil_moora.png",
            "seq_num": 9,
            "cap_text": "Implementasi Tabel Hasil MOORA (tbl_hasil_moora)"
        },
        {
            "num": "5.3.10",
            "title": "Implementasi Tabel Detail MOORA (tbl_detail_moora)",
            "desc": "Tabel Detail MOORA (tbl_detail_moora) berfungsi untuk menyimpan riwayat jejak tahapan kalkulasi matriks MOORA pada setiap kriteria penilaian, meliputi nilai matriks awal (X), nilai matriks ternormalisasi (X*), bobot preferensi kriteria (W), serta nilai matriks terbobot (W × X*). Implementasi data pada tabel Detail MOORA disajikan pada Gambar 5.10.",
            "img": "db_screenshots/img_db_10_detail_moora.png",
            "seq_num": 10,
            "cap_text": "Implementasi Tabel Detail MOORA (tbl_detail_moora)"
        },
        {
            "num": "5.3.11",
            "title": "Implementasi Tabel Log Aktivitas (tbl_log_aktivitas)",
            "desc": "Tabel Log Aktivitas (tbl_log_aktivitas) digunakan untuk merekam jejak audit trail seluruh aktivitas dan transaksi penting yang dilakukan pengguna dalam sistem, mencakup pencatatan user pengakses, modul yang digunakan, deskripsi tindakan, waktu pencatatan (timestamp), dan alamat IP pengguna. Implementasi data pada tabel Log Aktivitas disajikan pada Gambar 5.11.",
            "img": "db_screenshots/img_db_11_log_aktivitas.png",
            "seq_num": 11,
            "cap_text": "Implementasi Tabel Log Aktivitas (tbl_log_aktivitas)"
        }
    ]

    for item in db_sections:
        add_h3(item["num"], item["title"])
        add_p(item["desc"])
        add_image(item["img"], width_cm=15.2)
        p_cap = new_doc.add_paragraph()
        add_caption_field(p_cap, "Gambar 5. ", item["seq_num"], item["cap_text"], seq_type="Gambar")

    # 5.4 Implementasi Sistem
    add_h2("Implementasi Sistem")
    
    # 5.5 Pengujian Sistem
    add_h2("Pengujian Sistem")

    # Save to BAB V.docx
    new_doc.save("BAB V.docx")
    print("BAB V.docx successfully updated with Section 5.3 and all 11 database screenshots!")

if __name__ == "__main__":
    create_full_bab5_doc()
