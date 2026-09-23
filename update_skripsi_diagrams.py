import os
import docx
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement, parse_xml

# 15 Activity Diagram Descriptions (ringkas 3-4 baris)
ad_descriptions = {
    "AD-01": (
        "Activity Diagram Login menggambarkan alur autentikasi pengguna ke dalam sistem. "
        "Pengguna membuka halaman login, memasukkan username dan password, kemudian sistem memvalidasi data login. "
        "Jika data valid, sistem mengarahkan pengguna ke dashboard sesuai hak akses; jika tidak valid, sistem menampilkan pesan kesalahan."
    ),
    "AD-02": (
        "Activity Diagram Mengelola Data Pengguna menggambarkan alur pengelolaan akun pengguna oleh Administrator. "
        "Administrator membuka menu pengguna, memilih tindakan (tambah, ubah, atau hapus akun), lalu mengisi data yang diperlukan. "
        "Sistem memvalidasi kelengkapan data dan menyimpan perubahan ke basis data apabila data telah valid."
    ),
    "AD-03": (
        "Activity Diagram Mengelola Data Guru menggambarkan alur pengelolaan data master guru sebagai alternatif penilaian. "
        "Pengguna membuka menu data guru, menambahkan atau memperbarui profil guru (NIP, nama, mata pelajaran, beban mengajar). "
        "Sistem memvalidasi format data dan menyimpan perubahan ke basis data apabila seluruh isian telah lengkap dan valid."
    ),
    "AD-04": (
        "Activity Diagram Mengelola Kriteria dan Bobot menggambarkan alur konfigurasi kriteria penilaian kinerja guru. "
        "Administrator membuka menu kriteria, lalu mengubah nilai bobot preferensi atau sifat atribut kriteria (Benefit/Cost). "
        "Sistem memvalidasi bahwa total akumulasi bobot bernilai 1.00 (100%) sebelum menyimpan perubahan ke basis data."
    ),
    "AD-05": (
        "Activity Diagram Mengelola Skala Penilaian menggambarkan alur pengelolaan rubrik indikator skor penilaian (skala 1–5). "
        "Administrator memilih kriteria tertentu, kemudian menambahkan atau memperbarui deskripsi indikator ketercapaian. "
        "Sistem memvalidasi rentang nilai skor dan menyimpan perubahan rubrik ke basis data apabila data valid."
    ),
    "AD-06": (
        "Activity Diagram Mengelola Periode Penilaian menggambarkan alur manajemen siklus evaluasi berkala guru. "
        "Administrator membuka menu periode, membuat periode baru atau mengubah status periode (Draft, Aktif, Selesai). "
        "Sistem memastikan hanya satu periode berstatus Aktif dalam satu waktu, lalu menyimpan perubahan ke basis data."
    ),
    "AD-07": (
        "Activity Diagram Melakukan Penilaian Kinerja Guru menggambarkan alur pengisian skor evaluasi oleh Kepala Sekolah. "
        "Kepala Sekolah memilih periode aktif dan guru yang akan dinilai, kemudian memberikan skor pada kriteria C1 sampai C4. "
        "Sistem memvalidasi kelengkapan nilai dan menyimpan hasil transaksi penilaian ke basis data."
    ),
    "AD-08": (
        "Activity Diagram Mencatat Hasil Penilaian menggambarkan alur pencatatan dan verifikasi data evaluasi oleh Administrator. "
        "Administrator membuka rekapitulasi penilaian, memeriksa kelengkapan data nilai seluruh guru pada periode aktif. "
        "Sistem melakukan pengecekan integritas data dan mengonfirmasi pencatatan hasil penilaian apabila data telah lengkap."
    ),
    "AD-09": (
        "Activity Diagram Mengolah Penilaian dengan MOORA menggambarkan alur eksekusi algoritma SPK metode MOORA. "
        "Administrator menjalankan proses kalkulasi; sistem memeriksa kelengkapan data, lalu mengeksekusi tahapan MOORA: "
        "pembentukan Matriks Keputusan, Normalisasi, Pembobotan, perhitungan Nilai Preferensi Yi, dan penyimpanan hasil ranking ke basis data."
    ),
    "AD-10": (
        "Activity Diagram Melihat Pemeringkatan MOORA menggambarkan alur peninjauan hasil ranking kinerja guru. "
        "Pengguna membuka menu pemeringkatan dan memilih periode penilaian yang diinginkan. "
        "Sistem mengambil data preferensi Yi dan menyajikan tabel leaderboard ranking guru secara terurut dari peringkat tertinggi."
    ),
    "AD-11": (
        "Activity Diagram Melihat Hasil Penilaian menggambarkan alur peninjauan skor akhir evaluasi oleh pengguna. "
        "Pengguna mengakses menu hasil penilaian dan memilih periode evaluasi yang ingin dilihat. "
        "Sistem memvalidasi hak akses, mengambil data capaian kinerja, dan menyajikan ringkasan nilai beserta predikat kelayakan."
    ),
    "AD-12": (
        "Activity Diagram Melihat Detail Penilaian menggambarkan alur peninjauan rincian nilai per kriteria untuk setiap guru. "
        "Pengguna memilih rekaman hasil evaluasi dan mengklik tombol lihat detail. "
        "Sistem menampilkan rincian skor per kriteria (C1 Pedagogik, C2 Profesional, C3 Kepribadian, C4 Sosial) beserta catatan evaluator."
    ),
    "AD-13": (
        "Activity Diagram Monitoring Kinerja Guru menggambarkan alur pemantauan progres evaluasi guru secara berkala. "
        "Pengguna membuka modul monitoring kinerja dan memilih periode atau kategori guru yang ingin dipantau. "
        "Sistem menyajikan grafik statistik persentase guru yang sudah dinilai dan belum dinilai pada periode tersebut."
    ),
    "AD-14": (
        "Activity Diagram Melihat Riwayat Penilaian menggambarkan alur peninjauan rekam jejak evaluasi dari periode-periode sebelumnya. "
        "Pengguna membuka menu riwayat penilaian dan memilih periode tertentu. "
        "Sistem menyajikan riwayat skor, nilai preferensi MOORA, dan grafik perkembangan performa guru antar periode."
    ),
    "AD-15": (
        "Activity Diagram Mencetak Laporan Penilaian menggambarkan alur pencetakan dokumen hasil evaluasi kinerja guru. "
        "Pengguna membuka menu laporan, memilih periode evaluasi, dan menentukan format ekspor (PDF atau Excel). "
        "Sistem mengompilasi rekapitulasi nilai, ranking MOORA, dan lembar pengesahan, lalu menghasilkan dokumen laporan siap cetak."
    )
}

# 15 Sequence Diagram Descriptions (ringkas 3-4 baris)
sd_descriptions = {
    "SD-01": (
        "Sequence Diagram Login menguraikan urutan interaksi saat pengguna melakukan autentikasi ke sistem. "
        "Pengguna membuka Halaman Login, memasukkan kredensial, lalu sistem memeriksa data akun pada penyimpanan data. "
        "Pada fragmen alternatif, jika data valid sistem mengarahkan ke dashboard; jika tidak valid, sistem menampilkan pesan kesalahan."
    ),
    "SD-02": (
        "Sequence Diagram Mengelola Data Pengguna menguraikan interaksi pengelolaan akun pengguna oleh Administrator. "
        "Administrator membuka Halaman Data Pengguna, sistem menampilkan tabel akun terdaftar dari penyimpanan data. "
        "Administrator menginputkan perubahan data, sistem memvalidasi dan menyimpan ke basis data, lalu menampilkan pembaruan."
    ),
    "SD-03": (
        "Sequence Diagram Mengelola Data Guru menguraikan interaksi pengelolaan master data guru alternatif. "
        "Pengguna membuka Halaman Data Guru, sistem mengambil dan menampilkan daftar guru dari penyimpanan data. "
        "Pengguna menginput atau mengedit data guru, sistem memvalidasi format NIP dan kelengkapan data, lalu menyimpan perubahan."
    ),
    "SD-04": (
        "Sequence Diagram Mengelola Kriteria dan Bobot menguraikan interaksi pembaruan bobot preferensi kriteria MOORA. "
        "Administrator membuka Halaman Kriteria & Bobot, sistem menampilkan data kriteria dari penyimpanan data. "
        "Administrator mengubah nilai bobot, sistem memvalidasi total bobot bernilai 100%, lalu menyimpan dan memperbarui tampilan."
    ),
    "SD-05": (
        "Sequence Diagram Mengelola Skala Penilaian menguraikan interaksi konfigurasi rubrik indikator skor penilaian. "
        "Administrator membuka Halaman Skala Penilaian, sistem menampilkan tabel rubrik dari penyimpanan data. "
        "Administrator memperbarui deskripsi indikator skala 1–5, sistem memvalidasi dan menyimpan perubahan ke basis data."
    ),
    "SD-06": (
        "Sequence Diagram Mengelola Periode Penilaian menguraikan interaksi pengaturan siklus periode evaluasi. "
        "Administrator membuka Halaman Periode Penilaian, sistem menampilkan daftar periode dari penyimpanan data. "
        "Administrator membuat atau mengubah status periode, sistem memvalidasi keaktifan dan menyimpan perubahan ke basis data."
    ),
    "SD-07": (
        "Sequence Diagram Melakukan Penilaian Kinerja Guru menguraikan interaksi pengisian skor evaluasi oleh Kepala Sekolah. "
        "Kepala Sekolah membuka Halaman Penilaian, sistem menyajikan formulir evaluasi kriteria C1–C4 untuk guru yang dipilih. "
        "Kepala Sekolah menginput skor, sistem memvalidasi kelengkapan nilai dan menyimpan transaksi penilaian ke basis data."
    ),
    "SD-08": (
        "Sequence Diagram Mencatat Hasil Penilaian menguraikan interaksi pencatatan dan verifikasi data evaluasi. "
        "Administrator membuka Halaman Rekapitulasi Penilaian, sistem menampilkan daftar entri evaluasi guru dari penyimpanan data. "
        "Administrator mengonfirmasi rekapitulasi, sistem memverifikasi integritas data dan menyimpan status pencatatan."
    ),
    "SD-09": (
        "Sequence Diagram Mengolah Penilaian dengan MOORA menguraikan interaksi komputasi algoritma SPK metode MOORA. "
        "Administrator membuka Halaman Pengolahan MOORA dan menekan tombol proses kalkulasi. "
        "Sistem memvalidasi kelengkapan data, mengeksekusi tahapan MOORA (Matriks X, Normalisasi X*, Pembobotan V, Nilai Yi), dan menyimpan hasil ranking."
    ),
    "SD-10": (
        "Sequence Diagram Melihat Pemeringkatan MOORA menguraikan interaksi penayangan peringkat kinerja guru. "
        "Pengguna membuka Halaman Ranking MOORA dan memilih periode penilaian. "
        "Sistem mengambil data preferensi Yi dari penyimpanan Hasil MOORA dan menampilkan tabel ranking guru secara terurut."
    ),
    "SD-11": (
        "Sequence Diagram Melihat Hasil Penilaian menguraikan interaksi penayangan hasil evaluasi bagi seluruh pengguna. "
        "Pengguna membuka Halaman Hasil Penilaian, sistem memvalidasi hak akses dan mengambil data evaluasi dari penyimpanan. "
        "Sistem menyajikan ringkasan nilai, status evaluasi, dan predikat capaian kinerja guru yang bersangkutan."
    ),
    "SD-12": (
        "Sequence Diagram Melihat Detail Penilaian menguraikan interaksi peninjauan rincian nilai per kriteria. "
        "Pengguna memilih data guru dan mengklik tombol Lihat Detail, sistem mengambil rincian skor dari penyimpanan data. "
        "Sistem menampilkan breakdown skor C1 sampai C4 lengkap dengan catatan evaluasi dari penilai."
    ),
    "SD-13": (
        "Sequence Diagram Monitoring Kinerja Guru menguraikan interaksi pemantauan progres penilaian guru. "
        "Pengguna membuka Halaman Monitoring Kinerja, sistem mengambil data perbandingan guru yang sudah dan belum dinilai. "
        "Sistem menghitung persentase kelengkapan dan menyajikan grafik statistik progres penilaian secara visual."
    ),
    "SD-14": (
        "Sequence Diagram Melihat Riwayat Penilaian menguraikan interaksi peninjauan rekam jejak evaluasi antarperiode. "
        "Pengguna membuka Halaman Riwayat Penilaian, sistem mengambil arsip evaluasi dari penyimpanan Data Periode & Hasil MOORA. "
        "Sistem menampilkan histori penilaian beserta grafik tren performa kinerja guru dari periode ke periode."
    ),
    "SD-15": (
        "Sequence Diagram Mencetak Laporan Penilaian menguraikan interaksi pencetakan laporan resmi hasil evaluasi. "
        "Pengguna membuka Halaman Laporan dan memilih periode evaluasi serta format ekspor (PDF/Excel). "
        "Sistem mengompilasi data rekapitulasi nilai, ranking MOORA, dan lembar tanda tangan, lalu menghasilkan dokumen laporan siap unduh."
    )
}


def update_skripsi_document():
    src_docx = "SKRIPSI.docx"
    doc = docx.Document(src_docx)
    
    print(f"Total initial paragraphs in SKRIPSI.docx: {len(doc.paragraphs)}")
    
    # We will iterate through paragraphs and insert descriptions after the image paragraphs of AD-01..AD-15 and SD-01..SD-15
    # Notice: In docx, to insert a paragraph after a specific paragraph element, we can insert into _p parent XML
    
    def insert_paragraph_after(target_paragraph, text, bold_prefix=""):
        new_p = OxmlElement('w:p')
        target_paragraph._p.addnext(new_p)
        new_para = docx.text.paragraph.Paragraph(new_p, target_paragraph._parent)
        
        # Style formatting
        new_para.paragraph_format.line_spacing = 1.5
        new_para.paragraph_format.space_before = Pt(4)
        new_para.paragraph_format.space_after = Pt(8)
        new_para.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        
        if bold_prefix:
            r_pre = new_para.add_run(bold_prefix + " ")
            r_pre.font.name = 'Times New Roman'
            r_pre.font.size = Pt(12)
            r_pre.font.bold = True
            
        r_txt = new_para.add_run(text)
        r_txt.font.name = 'Times New Roman'
        r_txt.font.size = Pt(12)
        return new_para

    # ============================================================
    # STEP 1: Remove any existing "Penjelasan AD-" / "Penjelasan SD-" paragraphs
    # (from the previous run that placed them AFTER the image)
    # ============================================================
    removed = 0
    for p in list(doc.paragraphs):
        if p.text.strip().startswith("Penjelasan AD-") or p.text.strip().startswith("Penjelasan SD-"):
            p._p.getparent().remove(p._p)
            removed += 1
    print(f"[CLEANUP] Removed {removed} old explanation paragraphs")

    # ============================================================
    # STEP 2: Insert explanations BEFORE the image (after the title)
    # Structure target: Title (AD-xx) → Explanation → Image
    # ============================================================
    modified_count = 0
    paragraphs = doc.paragraphs  # refresh after removal

    for i in range(len(paragraphs)):
        p = paragraphs[i]
        txt = p.text.strip()

        # Check Activity Diagrams
        for ad_key, ad_desc in ad_descriptions.items():
            if txt.startswith(ad_key):
                # Insert AFTER the title paragraph p (i.e. BEFORE the image at i+1)
                insert_paragraph_after(p, ad_desc, bold_prefix=f"Penjelasan {ad_key}:")
                modified_count += 1
                print(f"[SUCCESS] Added explanation for {ad_key} (before image)")
                break

        # Check Sequence Diagrams
        for sd_key, sd_desc in sd_descriptions.items():
            if txt.startswith(sd_key):
                insert_paragraph_after(p, sd_desc, bold_prefix=f"Penjelasan {sd_key}:")
                modified_count += 1
                print(f"[SUCCESS] Added explanation for {sd_key} (before image)")
                break

    print(f"Total diagram explanations inserted: {modified_count}")

    # Save to SKRIPSI.docx
    out_docx = "SKRIPSI.docx"
    try:
        doc.save(out_docx)
        print(f"[SUCCESS] SKRIPSI.docx berhasil diperbarui secara langsung!")
    except PermissionError:
        alt_docx = "SKRIPSI_Updated.docx"
        doc.save(alt_docx)
        print(f"[WARNING] SKRIPSI.docx sedang dibuka di Microsoft Word.")
        print(f"[SUCCESS] Disimpan sebagai: {alt_docx}")

if __name__ == "__main__":
    update_skripsi_document()
