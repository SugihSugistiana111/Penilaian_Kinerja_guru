import docx

doc = docx.Document('BAB V.docx')
print(f"Total Paragraphs: {len(doc.paragraphs)}")
print(f"Total Tables: {len(doc.tables)}")

img_count = 0
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._r.xml for r in p.runs)
    if has_img:
        img_count += 1
        print(f"  [IMG #{img_count}] at P{i}")
    elif p.text.strip():
        print(f"  P{i} [{p.style.name}]: {p.text[:90]}...")

print(f"\nTotal Images found: {img_count}")
