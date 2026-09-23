import docx

doc = docx.Document('BAB V.docx')
print(f"Total Paragraphs in BAB V.docx: {len(doc.paragraphs)}")
print(f"Total Tables in BAB V.docx: {len(doc.tables)}")

img_count = 0
for i, p in enumerate(doc.paragraphs):
    has_img = any('graphic' in r._r.xml for r in p.runs)
    if has_img:
        img_count += 1
        print(f"  [IMG #{img_count}] at P{i}")
    elif p.text.strip():
        print(f"  P{i}: '{p.text[:95]}'")

print(f"\nTotal Figures/Images found: {img_count}")
